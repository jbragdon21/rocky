"""
Inbox Cleaner — per-user inbox triage at 200k+ message scale.

Each colleague gets a named process ("inbox-matt", "inbox-paul", ...) that
starts as a one-time deep clean and becomes a permanent, personalized
maintenance process. Design ground truth: BUILD_REFERENCE.md "Inbox Cleaner".

    python rocky.py --inbox-matt --snapshot [--full]
        Pull the full folder tree + inbox metadata (no bodies, no
        attachments) into a local JSONL snapshot. Checkpointed/resumable;
        re-runs fetch only mail newer than the cursor. Read-only:
        app-token Mail.Read via the Application Access Policy.

    python rocky.py --inbox-matt --analyze
        Build the sender + conversation ledgers, write the review workbook
        (ledgers + cohorts + drill-down), and generate draft cohorts:
        the user's MATTERS (from the share folder's matters.json, seeded
        from the questionnaire — subject-keyword match, whole conversations,
        Client\\Matter folders), standing sender routes (sender_routes.json),
        court-notice senders, newsletters/bulk (minus routed and never-bulk
        senders and anyone who writes on a known matter), old internal
        office traffic, and big case clusters (which need a folder label).
        Re-runs prune undecided drafts the current rules no longer produce,
        so editing matters.json/sender_routes.json + re-running --analyze
        regenerates cleanly. Templates: _templates\\inbox_matters.example.json
        and _templates\\inbox_sender_routes.example.json.

    python rocky.py --inbox-matt --questionnaire [--resend]
        Onboarding questionnaire BY EMAIL (not Teams — too many questions
        for chat back-and-forth). First run sends it from rocky@ as HTML
        with an answer box under each question ("Answer [iq-N]:" markers,
        same reply-parsing pattern as the Maple digest question boxes).
        Subsequent runs watch rocky@'s inbox for the user's reply, parse
        the typed answers out of the boxes, save them to the share folder,
        and feed them to the nightly rules update. Idempotent — schedule
        it alongside --chat until the answers arrive; it goes quiet after.

    python rocky.py --inbox-matt --chat
        One Teams poll cycle: log every new reply, resolve the pending
        cohort proposal (YES/NO/label), then propose the next draft cohort.
        Schedule every 15 minutes during an active cleanup.

    python rocky.py --inbox-matt --rules-update
        The daily Claude pass: fold the day's decisions, questionnaire
        answers, and chat traffic into the user's plain-English rules file
        (rules.md), then tune matters.json from the chat (structured
        add/update/close ops only, never delete; backed up + logged; the
        digest reports changes) and re-run analyze so the next chat cycle
        proposes the updated batches. Skips all API calls on quiet days.

    python rocky.py --inbox-matt --digest [--hours N]
        Daily plain-English digest of everything the process did in the
        window (default 24h), emailed from rocky@ to the mailbox owner and
        the observers (James). Built from the activity/communications/move
        logs; one Claude call phrases it plainly, with a deterministic
        fallback if the API is unavailable. Quiet window = no email.

    python rocky.py --inbox-matt --execute [--live] [--limit N]
        Execute approved cohorts. DRY-RUN by default — prints/logs the move
        plan. --live performs the moves. Auth per config write_via: "app"
        (application Mail.ReadWrite + Access Policy) or "delegated"
        (rocky@'s Mail.ReadWrite.Shared + an Exchange Full Access
        delegation on the mailbox — Matt's path). Everything is moved,
        never deleted; newsletters go to Deleted Items. Every move is
        logged with its source folder for undo-by-replay.

    python rocky.py --inbox-matt --status
        One-screen summary: snapshot size, cohort states, last chat.

Storage:
    local (C:\\Rocky\\inbox_cleaner\\<user>\\)  — snapshot.jsonl, folders.json,
        moves.jsonl (big / machine data)
    shared (OneDrive "Rocky Inboxes\\inbox-<user>\\") — rules.md, cohorts.json,
        activity.jsonl, communications.jsonl, review workbook, questionnaire
        (human-facing; James can read everything the process did and said)

Aggressive logging is a feature: every step -> activity.jsonl, every Teams
message in either direction -> communications.jsonl, everything mirrored to
rocky.log with an [inbox-<user>] tag.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests

log = logging.getLogger("rocky.inbox_cleaner")

GRAPH_API_BASE = "https://graph.microsoft.com/v1.0"
GRAPH_PAGE_SIZE = 100

CLAUDE_MODEL = "claude-sonnet-4-5"
RULES_MAX_TOKENS = 4096

INTERNAL_DOMAINS = {"gallagherllp.com", "gejlaw.com"}

# Sender-address markers for automated court / e-filing mail.
COURT_SENDER_MARKERS = (
    "uscourts.gov", "ecfnotice", "ecf_bounces", "cmecf", "efiling",
    "e-filing", "casefilexpress", "efile", "courts.state",
    "dccourts.gov", "mdcourts.gov", "vacourts.gov",
)

# Local-part markers suggesting bulk / newsletter senders.
BULK_LOCALPART_MARKERS = (
    "noreply", "no-reply", "no_reply", "donotreply", "do-not-reply",
    "newsletter", "news", "marketing", "updates", "notifications",
    "notification", "alerts", "alert", "mailer", "email", "info",
    "bounce", "campaigns", "reply",
)

# Per-user tunables (overridable in the user's inbox_users config block).
DEFAULT_AGE_FLOOR_DAYS = 180        # cases + internal office traffic
DEFAULT_BIG_CLUSTER_MIN = 150       # messages before a thread family is "a big case"
DEFAULT_NEWSLETTER_MIN_COUNT = 50   # per-sender volume floor for the bulk cohort
DEFAULT_NEWSLETTER_MAX_READ_RATE = 0.35
DEFAULT_INTERNAL_MIN_COUNT = 10     # per-sender floor for the internal cohort

_SUBJECT_PREFIX_RE = re.compile(
    r"^\s*(?:(?:re|fw|fwd)\s*:|\[(?:external|ext)\]\s*:?)\s*", re.IGNORECASE
)

_YES_WORDS = {"yes", "y", "ok", "okay", "approve", "approved", "yep", "yes.", "👍"}
_NO_WORDS = {"no", "n", "skip", "decline", "declined", "nope", "no."}


# =============================================================================
# Paths / state / logging
# =============================================================================

def user_paths(config: dict, data_dir: Path, user_key: str) -> dict:
    """All file locations for one user's process. Local = big/machine data;
    share (OneDrive) = everything a human should be able to read."""
    share_root = config.get("inbox_cleaner_dir")
    if share_root:
        share_root = Path(share_root)
    else:
        cases_root = Path(config.get(
            "cases_root",
            r"C:\Users\jbragdon\OneDrive\OneDrive - gejlaw.com\Rocky Cases"))
        share_root = cases_root.parent / "Rocky Inboxes"
    share = share_root / f"inbox-{user_key}"
    local = data_dir / "inbox_cleaner" / user_key
    return {
        "user_key": user_key,
        "tag": f"[inbox-{user_key}]",
        "local": local,
        "share": share,
        "snapshot": local / "snapshot.jsonl",
        "folders": local / "folders.json",
        "moves": local / "moves.jsonl",
        "state": data_dir / "state" / f"inbox_cleaner_{user_key}.json",
        "cohorts": share / "cohorts.json",
        "rules": share / "rules.md",
        "rules_history": share / "rules_history",
        "activity": share / "activity.jsonl",
        "comms": share / "communications.jsonl",
        "workbook": share / f"inbox-{user_key} review.xlsx",
        "questionnaire": share / "questionnaire.md",
        "answers": share / "questionnaire_answers.md",
        "matters": share / "matters.json",
        "routes": share / "sender_routes.json",
    }


def _now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _append_jsonl(path: Path, record: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")


def log_event(paths: dict, kind: str, **fields) -> None:
    """Aggressive step logging: structured event to activity.jsonl AND a
    human line to rocky.log."""
    record = {"ts": _now_iso(), "event": kind, **fields}
    _append_jsonl(paths["activity"], record)
    summary = ", ".join(f"{k}={v}" for k, v in fields.items()
                        if k not in ("detail",) and not isinstance(v, (list, dict)))
    log.info(f"{paths['tag']} {kind}: {summary}" if summary
             else f"{paths['tag']} {kind}")


def log_comm(paths: dict, direction: str, text: str, *,
             message_id: str | None = None, cohort_id: str | None = None,
             sender: str | None = None) -> None:
    """Every Teams message, both directions, verbatim."""
    _append_jsonl(paths["comms"], {
        "ts": _now_iso(), "direction": direction, "message_id": message_id,
        "cohort_id": cohort_id, "sender": sender, "text": text,
    })
    preview = text if len(text) <= 160 else text[:157] + "..."
    log.info(f"{paths['tag']} teams {direction}"
             + (f" (cohort {cohort_id})" if cohort_id else "") + f": {preview}")


def load_state(paths: dict) -> dict:
    p = paths["state"]
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def save_state(paths: dict, state: dict) -> None:
    p = paths["state"]
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(state, indent=2), encoding="utf-8")


# =============================================================================
# Graph plumbing (read side)
# =============================================================================

def _graph_get(url: str, token: str, params: dict | None = None,
               max_retries: int = 5, extra_headers: dict | None = None):
    """GET with throttling/backoff — the snapshot runs patiently for hours."""
    headers = {"Authorization": f"Bearer {token}", "Accept": "application/json"}
    if extra_headers:
        headers.update(extra_headers)
    delay = 2.0
    for attempt in range(max_retries):
        try:
            resp = requests.get(url, headers=headers, params=params, timeout=60)
        except requests.RequestException as e:
            log.warning(f"[inbox-cleaner] network error ({attempt + 1}/{max_retries}): {e}")
            time.sleep(delay)
            delay = min(delay * 2, 60)
            continue
        if resp.status_code == 429 or resp.status_code >= 500:
            retry_after = resp.headers.get("Retry-After")
            wait = float(retry_after) if retry_after else delay
            log.info(f"[inbox-cleaner] throttled ({resp.status_code}); sleeping {wait:.0f}s")
            time.sleep(wait)
            delay = min(delay * 2, 60)
            continue
        return resp
    return None


def _graph_post(url: str, token: str, payload: dict, max_retries: int = 5):
    headers = {"Authorization": f"Bearer {token}", "Accept": "application/json"}
    delay = 2.0
    for attempt in range(max_retries):
        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=60)
        except requests.RequestException as e:
            log.warning(f"[inbox-cleaner] network error ({attempt + 1}/{max_retries}): {e}")
            time.sleep(delay)
            delay = min(delay * 2, 60)
            continue
        if resp.status_code == 429 or resp.status_code >= 500:
            retry_after = resp.headers.get("Retry-After")
            wait = float(retry_after) if retry_after else delay
            log.info(f"[inbox-cleaner] throttled ({resp.status_code}); sleeping {wait:.0f}s")
            time.sleep(wait)
            delay = min(delay * 2, 60)
            continue
        return resp
    return None


# =============================================================================
# Stage 1 — snapshot
# =============================================================================

def fetch_folder_tree(token: str, mailbox: str) -> list[dict]:
    """Walk the whole mailbox folder tree. Returns flat list of
    {id, displayName, path, parentId, totalItemCount, unreadItemCount}."""
    out: list[dict] = []

    def walk(parent_id: str | None, prefix: str):
        if parent_id:
            url = f"{GRAPH_API_BASE}/users/{mailbox}/mailFolders/{parent_id}/childFolders"
        else:
            url = f"{GRAPH_API_BASE}/users/{mailbox}/mailFolders"
        params: dict | None = {
            "$top": "200", "includeHiddenFolders": "true",
            "$select": "id,displayName,parentFolderId,childFolderCount,"
                       "totalItemCount,unreadItemCount",
        }
        next_url: str | None = None
        page = 0
        while page == 0 or next_url:
            resp = _graph_get(next_url or url, token,
                              None if next_url else params)
            if resp is None or resp.status_code != 200:
                log.error(f"[inbox-cleaner] folder walk failed at {prefix or '<root>'}: "
                          f"{resp.status_code if resp is not None else 'no response'}")
                return
            data = resp.json()
            for f in data.get("value", []):
                path = (prefix + "\\" if prefix else "") + (f.get("displayName") or "?")
                out.append({
                    "id": f.get("id"), "displayName": f.get("displayName"),
                    "path": path, "parentId": f.get("parentFolderId"),
                    "totalItemCount": f.get("totalItemCount"),
                    "unreadItemCount": f.get("unreadItemCount"),
                })
                if (f.get("childFolderCount") or 0) > 0:
                    walk(f["id"], path)
            next_url = data.get("@odata.nextLink")
            page += 1

    walk(None, "")
    return out


def _addr_of(obj: dict | None) -> tuple[str, str]:
    ea = ((obj or {}).get("emailAddress") or {})
    return ((ea.get("address") or "").strip().lower(),
            (ea.get("name") or "").strip())


def _slim(msg: dict) -> dict:
    """Reduce a Graph message to the metadata the funnel needs (~300 bytes)."""
    from_addr, from_name = _addr_of(msg.get("from"))
    to = [_addr_of(r)[0] for r in (msg.get("toRecipients") or [])]
    cc = [_addr_of(r)[0] for r in (msg.get("ccRecipients") or [])]
    return {
        "id": msg.get("id"),
        "conv": msg.get("conversationId"),
        "subject": (msg.get("subject") or "")[:300],
        "from": from_addr,
        "from_name": from_name[:80],
        "to": [a for a in to if a][:10],
        "cc": [a for a in cc if a][:10],
        "received": msg.get("receivedDateTime"),
        "read": bool(msg.get("isRead")),
        "attach": bool(msg.get("hasAttachments")),
        "importance": msg.get("importance"),
    }


def snapshot(token: str, config: dict, ucfg: dict, paths: dict,
             full: bool = False) -> dict:
    """Stage 1: folder tree + inbox metadata into local files. Cursor-based
    and resumable — safe to re-run any time; only newer mail is fetched."""
    mailbox = ucfg["mailbox"]
    state = load_state(paths)
    log_event(paths, "snapshot_start", mailbox=mailbox, full=full)

    folders = fetch_folder_tree(token, mailbox)
    paths["folders"].parent.mkdir(parents=True, exist_ok=True)
    paths["folders"].write_text(
        json.dumps({"fetched_at": _now_iso(), "folders": folders},
                   indent=1, ensure_ascii=False), encoding="utf-8")
    log_event(paths, "folder_tree_saved", folder_count=len(folders))

    if full and paths["snapshot"].exists():
        paths["snapshot"].unlink()
        state.pop("snapshot_cursor", None)
        log_event(paths, "snapshot_reset")

    cursor = state.get("snapshot_cursor")
    url = f"{GRAPH_API_BASE}/users/{mailbox}/mailFolders/inbox/messages"
    params: dict = {
        "$orderby": "receivedDateTime asc",
        "$top": str(GRAPH_PAGE_SIZE),
        "$select": "id,conversationId,subject,from,toRecipients,ccRecipients,"
                   "receivedDateTime,isRead,hasAttachments,importance",
    }
    if cursor:
        params["$filter"] = f"receivedDateTime gt {cursor}"

    fetched = 0
    max_received = cursor
    next_url: str | None = None
    page = 0
    token_refreshed = time.monotonic()
    while page == 0 or next_url:
        resp = _graph_get(next_url or url, token, None if next_url else params)
        if resp is None:
            log_event(paths, "snapshot_aborted", reason="network", fetched=fetched)
            break
        if resp.status_code == 401 and \
                time.monotonic() - token_refreshed > 120:
            # App tokens live ~1 hour; a big-mailbox snapshot outlives them
            # (seen 2026-07-09: died at 139k messages). Re-acquire and retry
            # the same page. A 401 on a *fresh* token is a real permission
            # problem and falls through to the abort below.
            from rocky import acquire_app_token  # lazy
            token = acquire_app_token(config)
            token_refreshed = time.monotonic()
            log_event(paths, "snapshot_token_refreshed", fetched=fetched)
            continue
        if resp.status_code != 200:
            log_event(paths, "snapshot_aborted", reason=f"http_{resp.status_code}",
                      fetched=fetched, detail=resp.text[:300])
            break
        data = resp.json()
        for msg in data.get("value", []):
            rec = _slim(msg)
            _append_jsonl(paths["snapshot"], rec)
            fetched += 1
            recv = rec.get("received")
            if recv and (max_received is None or recv > max_received):
                max_received = recv
        # Checkpoint the cursor every page so an interrupted run resumes.
        if max_received:
            state["snapshot_cursor"] = max_received
            save_state(paths, state)
        if fetched and fetched % 5000 == 0:
            log_event(paths, "snapshot_progress", fetched=fetched,
                      through=max_received)
        next_url = data.get("@odata.nextLink")
        page += 1

    log_event(paths, "snapshot_done", fetched_this_run=fetched,
              cursor=state.get("snapshot_cursor"))
    return {"fetched": fetched, "folders": len(folders)}


def load_snapshot(paths: dict) -> list[dict]:
    """Snapshot records, deduped by message id (cursor boundaries can
    double-fetch same-second messages)."""
    if not paths["snapshot"].exists():
        return []
    seen: set[str] = set()
    out: list[dict] = []
    with open(paths["snapshot"], "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            mid = rec.get("id")
            if not mid or mid in seen:
                continue
            seen.add(mid)
            out.append(rec)
    return out


# =============================================================================
# Stage 2/3 — ledgers, cohorts, workbook
# =============================================================================

def _domain(addr: str) -> str:
    return addr.rsplit("@", 1)[-1] if "@" in addr else ""


def _domain_in(addr: str, domains) -> bool:
    """True if the sender's domain is one of `domains` OR a subdomain of
    one (emails.woodmac.com matches woodmac.com — bulk senders routinely
    send from subdomains)."""
    d = _domain(addr)
    return any(d == dom or d.endswith("." + dom) for dom in domains)


def _is_internal(addr: str) -> bool:
    return _domain(addr) in INTERNAL_DOMAINS


def _looks_court(addr: str) -> bool:
    # "efile" is a substring of "sharefile" — ShareFile is a file-sharing
    # service, not a court (caught drafting Matt's cohorts, 2026-07-09).
    if "sharefile" in addr:
        return False
    return any(m in addr for m in COURT_SENDER_MARKERS)


def _looks_bulk_localpart(addr: str) -> bool:
    local = addr.split("@", 1)[0]
    return any(m in local for m in BULK_LOCALPART_MARKERS)


def normalize_subject(subject: str) -> str:
    """Strip RE:/FW:/[EXTERNAL] prefixes (ported from the SortByConversation
    VBA idea) so subject-family fallback matching works."""
    s = subject or ""
    for _ in range(6):
        new = _SUBJECT_PREFIX_RE.sub("", s)
        if new == s:
            break
        s = new
    return s.strip().lower()


# =============================================================================
# Per-user matter + sender-route files (seeded from the questionnaire)
# =============================================================================
# Two human-editable JSON files in the user's share folder drive the
# user-specific passes. Both are optional — a user with neither file gets
# only the generic passes (big clusters, court, newsletters, internal).
#
# matters.json — the user's deals/cases, by client. The matter pass matches
# message subjects against each matter's keywords (whole-word, case-
# insensitive) and claims the ENTIRE conversation when any message in it
# matches, so file-share/DocuSign notices about a deal ride along. Closed
# matters get no age floor (archive wholesale); open ones get the user's
# age floor. This is the transactional-practice answer to the litigation
# "big thread family" heuristic — deal mail is many small threads that
# share matter names in their subjects (seen 2026-07-09: Matt's 198k
# messages = 107k conversations, zero families over 150).
#
# sender_routes.json — standing sender→folder routes from the questionnaire
# (e.g. SEIA → "SEIA Folder"), plus never_bulk_domains: senders that must
# never be swept to Deleted Items by the newsletters pass (file-share /
# data-room / e-signature notices that belong with their deals).

_KIND_PRIORITY = {  # execute-time claim order: first claimant wins a message
    "matter": 0, "big_case": 1, "sender_route": 2,
    "court_notices": 3, "newsletters": 4,
    "internal_ops": 5, "internal_broadcast": 6, "internal_office": 7,
    "internal_person": 8, "internal_misc": 9,
}

# Chat proposal order (lower proposes first; ties broken by count desc).
# Most-obviously-non-work batches go to the user first (Matt's ask,
# 2026-07-12): ops mailboxes and firm-wide blasts before anything that
# could plausibly hide misfiled deal/case mail (per-person internal mail
# proposes LAST).
_PROPOSE_PRIORITY = {
    "internal_ops": 1, "internal_broadcast": 2, "newsletters": 3,
    "court_notices": 4, "sender_route": 5, "matter": 6, "big_case": 6,
    "internal_office": 8, "internal_person": 8, "internal_misc": 9,
}

# The internal sweep is split into reviewable subgroups (see propose_cohorts
# pass 4). Functional firm mailboxes — the most obviously non-work senders.
_OPS_LOCALPART_MARKERS = (
    "billing", "accountspayable", "payable", "payroll", "administrator",
    "switchboard", "postmaster", "noreply", "no-reply", "no_reply",
    "donotreply", "clientmatters", "marketing", "reception", "facilities",
    "helpdesk", "itsupport",
)
_OPS_LOCALPART_EXACT = {"info", "admin", "hr", "it", "office", "accounting",
                        "events", "library", "records"}
DEFAULT_INTERNAL_PERSON_MIN = 150     # msgs before a colleague gets their own batch
DEFAULT_BROADCAST_MIN_RECIPIENTS = 8  # visible To/Cc count that says "blast"


def _looks_ops_localpart(addr: str) -> bool:
    local = addr.split("@", 1)[0].lower()
    return local in _OPS_LOCALPART_EXACT or \
        any(m in local for m in _OPS_LOCALPART_MARKERS)


def _matter_noun(paths: dict, ucfg: dict) -> str:
    """The user-facing word for their unit of work: 'deal' for a
    transactional practice (Matt), 'case' for litigation. Read from the
    user's config (matter_noun) or matters.json ("noun"); default 'case'.
    Flows into cohort titles/descriptions, Teams proposals, and digests."""
    n = (ucfg.get("matter_noun") or "").strip()
    if n:
        return n
    p = paths.get("matters")
    if p and p.exists():
        try:
            n = (json.loads(p.read_text(encoding="utf-8"))
                 .get("noun") or "").strip()
        except (json.JSONDecodeError, OSError):
            n = ""
    return n or "case"


def load_matters(paths: dict) -> list[dict]:
    """Flat list of {client, matter, keywords, closed, folder} from
    matters.json. Missing/invalid file = no matter pass."""
    p = paths.get("matters")
    if not p or not p.exists():
        return []
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        log.warning(f"[inbox-cleaner] matters.json unreadable, skipping "
                    f"matter pass: {e}")
        return []
    out = []
    for cl in data.get("clients") or []:
        client = (cl.get("client") or "").strip()
        for m in cl.get("matters") or []:
            name = (m.get("matter") or "").strip()
            kws = [k.strip() for k in (m.get("keywords") or []) if k.strip()]
            if not (client and name and kws):
                continue
            out.append({
                "client": client, "matter": name, "keywords": kws,
                "closed": bool(m.get("closed")),
                "folder": (m.get("folder") or f"{client}\\{name}").strip(),
                "participant_domains": [
                    d.strip().lower()
                    for d in (m.get("participant_domains") or []) if d.strip()],
            })
    return out


def load_routes(paths: dict) -> tuple[list[dict], set[str]]:
    """(routes, never_bulk_domains) from sender_routes.json."""
    p = paths.get("routes")
    if not p or not p.exists():
        return [], set()
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        log.warning(f"[inbox-cleaner] sender_routes.json unreadable, "
                    f"skipping routes: {e}")
        return [], set()
    routes = []
    for r in data.get("routes") or []:
        title = (r.get("title") or "").strip()
        folder = (r.get("target_folder") or "").strip()
        domains = [d.strip().lower() for d in (r.get("domains") or []) if d.strip()]
        senders = [s.strip().lower() for s in (r.get("senders") or []) if s.strip()]
        if title and folder and (domains or senders):
            routes.append({"title": title, "target_folder": folder,
                           "domains": domains, "senders": senders,
                           "age_floor_days": int(r.get("age_floor_days") or 0)})
    never_bulk = {d.strip().lower()
                  for d in (data.get("never_bulk_domains") or []) if d.strip()}
    return routes, never_bulk


def load_chat_exclusions(paths: dict) -> tuple[set[str], set[str]]:
    """(senders, domains) Rocky must NEVER propose moves for — hand-edited
    into sender_routes.json (exclude_senders / exclude_domains). Honored by
    every proposal pass and at execute time."""
    p = paths.get("routes")
    if not p or not p.exists():
        return set(), set()
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return set(), set()
    return (
        {s.strip().lower() for s in (data.get("exclude_senders") or [])
         if isinstance(s, str) and s.strip()},
        {d.strip().lower().lstrip("@") for d in (data.get("exclude_domains") or [])
         if isinstance(d, str) and d.strip()},
    )


def _drop_excluded(records: list[dict], excl_senders: set[str],
                   excl_domains: set[str]) -> list[dict]:
    if not excl_senders and not excl_domains:
        return records
    return [r for r in records
            if (r.get("from") or "") not in excl_senders
            and not _domain_in(r.get("from") or "", excl_domains)]


def _conv_key_of(r: dict) -> str:
    return r.get("conv") or f"subj:{normalize_subject(r.get('subject') or '')}"


def _matter_conv_keys(records: list[dict], keywords: list[str],
                      participant_domains: list[str] | None = None) -> set[str]:
    """Conversation keys where any message subject contains any keyword
    (whole-word, case-insensitive).

    When participant_domains is set, a keyword hit alone is not enough:
    the conversation must ALSO involve (from/to/cc) someone at one of
    those domains, in any of its messages. This is the corroboration
    guard for generic keywords ('Twelve', 'Dimension') that would
    otherwise sweep unrelated mail."""
    pats = [re.compile(r"(?<![a-z0-9])" + re.escape(k.lower()) + r"(?![a-z0-9])")
            for k in keywords]
    keys: set[str] = set()
    for r in records:
        subj = normalize_subject(r.get("subject") or "")
        if subj and any(p.search(subj) for p in pats):
            keys.add(_conv_key_of(r))
    if not keys or not participant_domains:
        return keys
    evidenced: set[str] = set()
    for r in records:
        key = _conv_key_of(r)
        if key not in keys or key in evidenced:
            continue
        if any(_domain_in(a, participant_domains) for a in
               ([r.get("from") or ""] + (r.get("to") or []) + (r.get("cc") or []))
               if a):
            evidenced.add(key)
    return evidenced


def build_ledgers(records: list[dict]) -> tuple[dict, dict]:
    """Sender ledger + conversation ledger from the snapshot."""
    senders: dict[str, dict] = {}
    convs: dict[str, dict] = {}
    for r in records:
        addr = r.get("from") or "(unknown)"
        s = senders.setdefault(addr, {
            "addr": addr, "name": r.get("from_name") or "", "count": 0,
            "unread": 0, "first": None, "last": None,
            "internal": _is_internal(addr),
        })
        s["count"] += 1
        if not r.get("read"):
            s["unread"] += 1
        recv = r.get("received")
        if recv:
            if s["first"] is None or recv < s["first"]:
                s["first"] = recv
            if s["last"] is None or recv > s["last"]:
                s["last"] = recv

        # Conversation family: conversationId primary, normalized subject fallback.
        key = r.get("conv") or f"subj:{normalize_subject(r.get('subject') or '')}"
        c = convs.setdefault(key, {
            "key": key, "subject": normalize_subject(r.get("subject") or ""),
            "count": 0, "first": None, "last": None, "senders": {},
        })
        c["count"] += 1
        c["senders"][addr] = c["senders"].get(addr, 0) + 1
        if recv:
            if c["first"] is None or recv < c["first"]:
                c["first"] = recv
            if c["last"] is None or recv > c["last"]:
                c["last"] = recv
    return senders, convs


def _read_rate(s: dict) -> float:
    return 1.0 - (s["unread"] / s["count"]) if s["count"] else 0.0


def _match_key(kind: str, match: dict) -> str:
    blob = json.dumps({"kind": kind, "match": match}, sort_keys=True)
    return hashlib.md5(blob.encode("utf-8")).hexdigest()[:12]


def propose_cohorts(records: list[dict], senders: dict, convs: dict,
                    ucfg: dict, paths: dict) -> list[dict]:
    """Draft cohorts from the deterministic funnel passes. Every cohort
    still needs human approval via Teams (or the workbook) before anything
    moves. Conservative on purpose: unmatched mail is simply left alone."""
    age_floor = int(ucfg.get("age_floor_days", DEFAULT_AGE_FLOOR_DAYS))
    big_min = int(ucfg.get("big_cluster_min", DEFAULT_BIG_CLUSTER_MIN))
    news_min = int(ucfg.get("newsletter_min_count", DEFAULT_NEWSLETTER_MIN_COUNT))
    news_read = float(ucfg.get("newsletter_max_read_rate",
                               DEFAULT_NEWSLETTER_MAX_READ_RATE))
    internal_min = int(ucfg.get("internal_min_count", DEFAULT_INTERNAL_MIN_COUNT))

    drafts: list[dict] = []
    cutoff = (datetime.now(timezone.utc) - timedelta(days=age_floor)) \
        .strftime("%Y-%m-%dT%H:%M:%SZ")

    # Chat-learned exclusions ("skip emails from x"): no pass ever drafts
    # a move for these senders/domains.
    records = _drop_excluded(records, *load_chat_exclusions(paths))

    # --- 0. Matters (from matters.json, seeded by the questionnaire). ------
    # One cohort per matter; matches whole conversations whose subjects name
    # the matter. Closed matters archive wholesale (no age floor). All
    # matter-claimed conversations are excluded from every later pass.
    # User-facing wording uses the user's own noun ("deal" for Matt's
    # transactional practice, "case" for litigation).
    noun = _matter_noun(paths, ucfg)
    matters = load_matters(paths)
    matter_claimed_keys: set[str] = set()
    for m in matters:
        keys = _matter_conv_keys(records, m["keywords"],
                                 m["participant_domains"]) - matter_claimed_keys
        floor = 0 if m["closed"] else age_floor
        m_cutoff = (datetime.now(timezone.utc) - timedelta(days=floor)) \
            .strftime("%Y-%m-%dT%H:%M:%SZ") if floor else None
        matched = [r for r in records if _conv_key_of(r) in keys
                   and (not m_cutoff or (r.get("received") or "") < m_cutoff)]
        n = len(matched)
        if not n:
            continue
        matter_claimed_keys |= keys
        by_sender: dict[str, int] = {}
        for r in matched:
            frm = r.get("from") or "(unknown)"
            by_sender[frm] = by_sender.get(frm, 0) + 1
        state = f"closed {noun} — archive wholesale" if m["closed"] \
            else f"open {noun} — only mail older than {floor} days moves"
        guard = (f" Only conversations involving {', '.join(m['participant_domains'])} count."
                 if m["participant_domains"] else "")
        match: dict = {"client": m["client"], "matter": m["matter"],
                       "matter_keywords": m["keywords"]}
        if m["participant_domains"]:
            match["participant_domains"] = m["participant_domains"]
        drafts.append({
            "kind": "matter",
            "title": f"{m['client']} — {m['matter']}",
            "description": (
                f"{n:,} emails in conversations naming this {noun} "
                f"(keywords: {', '.join(m['keywords'])}). {state}.{guard}"),
            "match": match,
            "count": n,
            "top_senders": [f"{a} ({c})" for a, c in
                            sorted(by_sender.items(), key=lambda kv: -kv[1])[:5]],
            "target_folder": m["folder"],
            "disposition": "move",
            "age_floor_days": floor,
        })

    routes, never_bulk = load_routes(paths)
    route_domains = {d for r in routes for d in r["domains"]}
    route_senders = {s for r in routes for s in r["senders"]}

    big_conv_keys = {k for k, c in convs.items()
                     if c["count"] >= big_min and k not in matter_claimed_keys}

    # --- 1. Big case clusters: one label-cohort per thread family. --------
    for key in sorted(big_conv_keys, key=lambda k: -convs[k]["count"]):
        c = convs[key]
        top_senders = sorted(c["senders"].items(), key=lambda kv: -kv[1])[:5]
        drafts.append({
            "kind": "big_case",
            "title": f"Thread family: {c['subject'][:80] or '(no subject)'}",
            "description": (
                f"{c['count']:,} emails in one thread family "
                f"({(c['first'] or '?')[:10]} – {(c['last'] or '?')[:10]}). "
                f"Looks like a long-running {noun}."),
            "match": {"conversation_keys": [key]},
            "count": c["count"],
            "top_senders": [f"{a} ({n})" for a, n in top_senders],
            "target_folder": None,          # the user names it via Teams
            "disposition": "move",
            "age_floor_days": age_floor,
        })

    # --- 2. Court-notice senders (age floor applies). ----------------------
    court_senders = [a for a in senders if _looks_court(a)]
    if court_senders:
        n = sum(1 for r in records
                if r.get("from") in court_senders
                and (r.get("received") or "") < cutoff)
        if n:
            drafts.append({
                "kind": "court_notices",
                "title": "Automated court / e-filing notices",
                "description": (
                    f"{n:,} emails older than {age_floor} days from "
                    f"{len(court_senders)} court/e-filing sender(s)."),
                "match": {"senders": sorted(court_senders)},
                "count": n,
                "target_folder": "Court Notices Archive",
                "disposition": "move",
                "age_floor_days": age_floor,
            })

    # --- 2b. Standing sender routes (from sender_routes.json). -------------
    for rt in routes:
        rt_cutoff = (datetime.now(timezone.utc)
                     - timedelta(days=rt["age_floor_days"])) \
            .strftime("%Y-%m-%dT%H:%M:%SZ") if rt["age_floor_days"] else None
        n = sum(1 for r in records
                if (_domain_in(r.get("from") or "", rt["domains"])
                    or (r.get("from") or "") in rt["senders"])
                and (not rt_cutoff or (r.get("received") or "") < rt_cutoff))
        if not n:
            continue
        drafts.append({
            "kind": "sender_route",
            "title": rt["title"],
            "description": (
                f"{n:,} emails from standing-rule senders "
                f"({', '.join(rt['domains'] + rt['senders'])}) → "
                f"\"{rt['target_folder']}\"."),
            "match": {"domains": rt["domains"], "senders": rt["senders"]},
            "count": n,
            "target_folder": rt["target_folder"],
            "disposition": "move",
            "age_floor_days": rt["age_floor_days"],
        })

    # --- 3. Newsletters / bulk (no age floor). -----------------------------
    # Routed senders go to their own folders above; never_bulk domains
    # (file-share / data-room / e-signature notices) are left alone rather
    # than swept to Deleted Items — the matter pass files the ones whose
    # subjects name a deal. Anyone who writes on a known matter is a deal
    # contact, never a newsletter, no matter how bulk-ish their ledger looks.
    matter_senders = {r.get("from") for r in records
                      if _conv_key_of(r) in matter_claimed_keys}
    bulk = [
        a for a, s in senders.items()
        if not s["internal"] and not _looks_court(a)
        and a not in matter_senders
        and s["count"] >= news_min
        and (_looks_bulk_localpart(a) or _read_rate(s) < news_read)
        and not _domain_in(a, never_bulk)
        and not _domain_in(a, route_domains)
        and a not in route_senders
    ]
    if bulk:
        n = sum(1 for r in records if r.get("from") in bulk)
        top = sorted(bulk, key=lambda a: -senders[a]["count"])[:10]
        drafts.append({
            "kind": "newsletters",
            "title": "Newsletters & automated bulk mail",
            "description": (
                f"{n:,} emails from {len(bulk)} high-volume bulk sender(s) "
                f"(mostly unread / no-reply style). Plan: move to Deleted "
                f"Items — recoverable until retention clears them."),
            "match": {"senders": sorted(bulk)},
            "count": n,
            "top_senders": [f"{a} ({senders[a]['count']})" for a in top],
            "target_folder": "Deleted Items",
            "disposition": "deleteditems",
            "age_floor_days": 0,
        })

    # --- 4. Old, purely-internal office traffic — reviewable SUBGROUPS. ----
    # One 51k "internal" batch proved unreviewable (Matt, 2026-07-12: a deal
    # email with a colleague could hide in it). Split it: functional ops
    # mailboxes and firm-wide broadcasts are obviously non-work and propose
    # FIRST (_PROPOSE_PRIORITY); person-to-person mail comes as per-colleague
    # batches with sample subjects, proposed LAST, so misfiled deal/case
    # traffic can be spotted and re-routed before anything is approved.
    internal_senders = [
        a for a, s in senders.items()
        if s["internal"] and s["count"] >= internal_min
    ]

    def _purely_internal(r: dict) -> bool:
        return all(_is_internal(a) for a in
                   ([r.get("from") or ""] + (r.get("to") or []) + (r.get("cc") or []))
                   if a)

    protected_keys = big_conv_keys | matter_claimed_keys
    person_min = int(ucfg.get("internal_person_min",
                              DEFAULT_INTERNAL_PERSON_MIN))
    bcast_min = int(ucfg.get("broadcast_min_recipients",
                             DEFAULT_BROADCAST_MIN_RECIPIENTS))

    def _recips(r: dict) -> int:
        return len(r.get("to") or []) + len(r.get("cc") or [])

    def _int_draft(kind: str, title: str, desc: str, match: dict,
                   msgs: list[dict], samples: bool = False) -> None:
        if not msgs:
            return
        by_sender: dict[str, int] = {}
        for r in msgs:
            frm = r.get("from") or "(unknown)"
            by_sender[frm] = by_sender.get(frm, 0) + 1
        d = {
            "kind": kind, "title": title, "description": desc,
            "match": {**match, "purely_internal": True},
            "count": len(msgs),
            "top_senders": [f"{a} ({c})" for a, c in
                            sorted(by_sender.items(),
                                   key=lambda kv: -kv[1])[:8]],
            "target_folder": "Office Misc. Archive",
            "disposition": "move",
            "age_floor_days": age_floor,
        }
        if samples:
            subj: dict[str, int] = {}
            for r in msgs:
                s = normalize_subject(r.get("subject") or "")
                if s:
                    subj[s] = subj.get(s, 0) + 1
            d["sample_subjects"] = [s for s, _ in
                                    sorted(subj.items(),
                                           key=lambda kv: -kv[1])[:5]]
        drafts.append(d)

    ops_senders = sorted(a for a in internal_senders
                         if _looks_ops_localpart(a))
    ops_set = set(ops_senders)
    nonops = sorted(a for a in internal_senders if a not in ops_set)

    sweepable = [r for r in records
                 if r.get("from") in set(internal_senders)
                 and (r.get("received") or "") < cutoff
                 and _conv_key_of(r) not in protected_keys
                 and _purely_internal(r)]
    ops_msgs = [r for r in sweepable if r.get("from") in ops_set]
    rest = [r for r in sweepable if r.get("from") not in ops_set]
    bc_msgs = [r for r in rest if _recips(r) >= bcast_min]
    person_pool = [r for r in rest if _recips(r) < bcast_min]

    _int_draft(
        "internal_ops", "Office operations mail",
        (f"{len(ops_msgs):,} emails from functional firm mailboxes "
         f"(billing, accounts payable, administrator, switchboard, "
         f"no-reply and the like) older than {age_floor} days."),
        {"senders": ops_senders}, ops_msgs)
    _int_draft(
        "internal_broadcast", "Firm-wide announcements & broadcasts",
        (f"{len(bc_msgs):,} internal emails sent to {bcast_min}+ people at "
         f"once (announcements, events, office news) older than "
         f"{age_floor} days."),
        {"senders": nonops, "min_recipients": bcast_min}, bc_msgs)

    # Person-to-person internal mail, one batch per colleague (a person =
    # one localpart across the legacy/new domains). This is where a
    # misfiled {noun} email would hide — sample subjects included, proposed
    # last.
    groups: dict[str, dict] = {}
    for r in person_pool:
        lp = (r.get("from") or "").split("@", 1)[0]
        g = groups.setdefault(lp, {"addrs": set(), "msgs": []})
        g["addrs"].add(r.get("from"))
        g["msgs"].append(r)
    small_msgs: list[dict] = []
    for lp, g in sorted(groups.items(), key=lambda kv: -len(kv[1]["msgs"])):
        if len(g["msgs"]) < person_min:
            small_msgs.extend(g["msgs"])
            continue
        top_addr = max(g["addrs"],
                       key=lambda a: senders.get(a, {}).get("count", 0))
        name = (senders.get(top_addr) or {}).get("name") or lp
        _int_draft(
            "internal_person", f"Internal mail from {name}",
            (f"{len(g['msgs']):,} person-to-person internal emails from "
             f"{name} older than {age_floor} days, in no known {noun} "
             f"thread. Check the sample subjects — if any of this is "
             f"{noun} mail, say so and I'll re-route it instead of "
             f"archiving it."),
            {"senders": sorted(g["addrs"]),
             "max_recipients": bcast_min - 1},
            g["msgs"], samples=True)
    small_senders = sorted({r.get("from") for r in small_msgs if r.get("from")})
    _int_draft(
        "internal_misc",
        "Internal mail — everyone else (small volumes)",
        (f"{len(small_msgs):,} internal emails from {len(small_senders)} "
         f"colleagues with smaller volumes, older than {age_floor} days, "
         f"in no known {noun} thread."),
        {"senders": small_senders, "max_recipients": bcast_min - 1},
        small_msgs, samples=True)

    # Claim order for overlapping matches at execute time: file/draft order
    # within a kind (e.g. the two Stag Moose loans share a generic keyword —
    # the one listed first in matters.json wins its conversations).
    for i, d in enumerate(drafts):
        d["seq"] = i

    return drafts


def load_cohorts(paths: dict) -> list[dict]:
    if not paths["cohorts"].exists():
        return []
    try:
        return json.loads(paths["cohorts"].read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []


def save_cohorts(paths: dict, cohorts: list[dict]) -> None:
    paths["cohorts"].parent.mkdir(parents=True, exist_ok=True)
    paths["cohorts"].write_text(
        json.dumps(cohorts, indent=1, ensure_ascii=False), encoding="utf-8")


def analyze(config: dict, ucfg: dict, paths: dict) -> dict:
    """Stage 2+3: ledgers, draft cohorts (merged into cohorts.json without
    clobbering decisions already made), review workbook."""
    records = load_snapshot(paths)
    if not records:
        log_event(paths, "analyze_skipped", reason="no snapshot — run --snapshot first")
        return {"error": "no snapshot"}
    log_event(paths, "analyze_start", messages=len(records))

    senders, convs = build_ledgers(records)
    drafts = propose_cohorts(records, senders, convs, ucfg, paths)

    cohorts = load_cohorts(paths)
    # Self-cleaning regeneration: drop UNDECIDED drafts that the current
    # rules/data no longer produce (stale after a matters.json edit, config
    # change, or code fix). Anything proposed or decided is never touched.
    computed_keys = {_match_key(d["kind"], d["match"]) for d in drafts}
    stale = [c for c in cohorts if c.get("status") == "draft"
             and c.get("match_key") not in computed_keys]
    if stale:
        cohorts = [c for c in cohorts if c not in stale]
        log_event(paths, "cohorts_pruned", count=len(stale),
                  ids=[c.get("id") for c in stale])
    by_key = {c.get("match_key"): c for c in cohorts}
    next_num = max((int(c["id"][1:]) for c in cohorts
                    if re.fullmatch(r"C\d{4}", c.get("id", ""))), default=0) + 1
    added = 0
    refreshed = 0
    for d in drafts:
        mk = _match_key(d["kind"], d["match"])
        existing = by_key.get(mk)
        if existing is not None:
            # Same rule as a prior run. If it's still an undecided draft,
            # refresh what the current data/exclusions say about it (count,
            # wording, claim order); proposed/decided cohorts are frozen.
            if existing.get("status") == "draft":
                for fld in ("title", "description", "count", "target_folder",
                            "disposition", "age_floor_days", "seq"):
                    if fld in d:
                        existing[fld] = d[fld]
                refreshed += 1
            continue
        d.update({"id": f"C{next_num:04d}", "match_key": mk, "status": "draft",
                  "created_at": _now_iso(), "teams": {}})
        cohorts.append(d)
        by_key[mk] = d
        next_num += 1
        added += 1
    if refreshed:
        log_event(paths, "cohorts_refreshed", count=refreshed)
    save_cohorts(paths, cohorts)

    write_workbook(paths, records, senders, convs, cohorts)
    covered = sum(c["count"] for c in cohorts if c.get("status") != "declined")
    log_event(paths, "analyze_done", messages=len(records),
              senders=len(senders), conversations=len(convs),
              cohorts_total=len(cohorts), cohorts_added=added,
              messages_covered_by_cohorts=covered)
    return {"messages": len(records), "senders": len(senders),
            "conversations": len(convs), "cohorts_added": added,
            "workbook": str(paths["workbook"])}


def write_workbook(paths: dict, records: list[dict], senders: dict,
                   convs: dict, cohorts: list[dict]) -> None:
    """Human review workbook: cohort rows + the two ledgers (capped)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Cohorts"
    headers = ["ID", "Status", "Kind", "Count", "Title", "Target folder",
               "Disposition", "Age floor (days)", "Description"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
    for c in sorted(cohorts, key=lambda x: -x.get("count", 0)):
        ws.append([c.get("id"), c.get("status"), c.get("kind"), c.get("count"),
                   c.get("title"), c.get("target_folder"), c.get("disposition"),
                   c.get("age_floor_days"), c.get("description")])

    ws2 = wb.create_sheet("Sender ledger")
    ws2.append(["Sender", "Name", "Count", "Unread", "Read %", "Internal",
                "First", "Last"])
    for cell in ws2[1]:
        cell.font = bold
    for s in sorted(senders.values(), key=lambda x: -x["count"])[:1000]:
        ws2.append([s["addr"], s["name"], s["count"], s["unread"],
                    round(_read_rate(s) * 100), "Y" if s["internal"] else "",
                    (s["first"] or "")[:10], (s["last"] or "")[:10]])

    ws3 = wb.create_sheet("Conversation ledger")
    ws3.append(["Subject family", "Count", "First", "Last", "Top senders"])
    for cell in ws3[1]:
        cell.font = bold
    for c in sorted(convs.values(), key=lambda x: -x["count"])[:500]:
        top = sorted(c["senders"].items(), key=lambda kv: -kv[1])[:3]
        ws3.append([c["subject"][:120], c["count"], (c["first"] or "")[:10],
                    (c["last"] or "")[:10],
                    "; ".join(f"{a} ({n})" for a, n in top)])

    paths["workbook"].parent.mkdir(parents=True, exist_ok=True)
    wb.save(paths["workbook"])
    log_event(paths, "workbook_written", path=str(paths["workbook"]))


# =============================================================================
# Rules file
# =============================================================================

_RULES_SEED = """# Inbox rules — {display_name} ({process})

Maintained by Rocky's Inbox Cleaner. Every rule records where it came from
(questionnaire answer or a cohort decision made over Teams). Edit freely —
this file is the source of truth for future maintenance runs, and it is the
long-term value of the process.

## Standing preferences (from questionnaire)

(none yet)

## Learned rules

(none yet)

## Do not propose again

(none yet)
"""


def ensure_rules_file(paths: dict, ucfg: dict) -> None:
    if paths["rules"].exists():
        return
    paths["rules"].parent.mkdir(parents=True, exist_ok=True)
    paths["rules"].write_text(
        _RULES_SEED.format(display_name=ucfg.get("display_name", paths["user_key"]),
                           process=f"inbox-{paths['user_key']}"),
        encoding="utf-8")
    log_event(paths, "rules_file_created", path=str(paths["rules"]))


def _append_rule(paths: dict, section: str, line: str) -> None:
    """Insert a bullet at the end of a `## section`, dropping the
    '(none yet)' placeholder if present."""
    text = paths["rules"].read_text(encoding="utf-8")
    lines = text.splitlines()
    try:
        start = next(i for i, ln in enumerate(lines)
                     if ln.strip().lower() == f"## {section}".lower())
    except StopIteration:
        lines += ["", f"## {section}", "", line]
        paths["rules"].write_text("\n".join(lines) + "\n", encoding="utf-8")
        return
    end = next((i for i in range(start + 1, len(lines))
                if lines[i].startswith("## ")), len(lines))
    block = [ln for ln in lines[start + 1:end] if ln.strip() != "(none yet)"]
    while block and not block[-1].strip():
        block.pop()
    new = lines[:start + 1] + [""] + block + [line, ""] + lines[end:]
    paths["rules"].write_text("\n".join(new) + "\n", encoding="utf-8")


def record_decision_rule(paths: dict, cohort: dict, decision: str) -> None:
    """Persist a cohort decision as a durable rule with provenance."""
    stamp = datetime.now().strftime("%Y-%m-%d")
    match = cohort.get("match", {})
    if "senders" in match:
        crit = f"mail from {len(match['senders'])} sender(s) incl. " \
               + ", ".join(match["senders"][:3]) + ("..." if len(match["senders"]) > 3 else "")
    elif "conversation_keys" in match:
        crit = f"thread family \"{cohort.get('title', '')[:60]}\""
    else:
        crit = cohort.get("title", "?")
    floor = cohort.get("age_floor_days") or 0
    floor_txt = f" older than {floor} days" if floor else ""
    if decision == "approved":
        _append_rule(paths, "Learned rules",
                     f"- [{cohort['id']}] [approved {stamp}] Move {crit}{floor_txt} "
                     f"→ \"{cohort.get('target_folder')}\".")
    else:
        _append_rule(paths, "Do not propose again",
                     f"- [{cohort['id']}] [declined {stamp}] Leave {crit} alone.")
    log_event(paths, "rule_recorded", cohort=cohort["id"], decision=decision)


# =============================================================================
# Onboarding questionnaire — by email, answer boxes under each question
# =============================================================================
# The questionnaire goes out as an HTML email from rocky@ with a bordered
# answer box under every question. Each box carries an "Answer [iq-N]:"
# marker; the user types inside the box and hits Reply. Rocky watches her
# own inbox for the reply and lifts the answers out with the same
# marker-scan approach the Maple digest uses for its client-question boxes
# (pma_tracker._parse_digest_answers).

_Q_SUBJECT_PREFIX = "Rocky Inbox Cleaner questionnaire"

_Q_ANSWER_MARKER_RE = re.compile(r"Answer\s*\[\s*iq-(\d+)\s*\]\s*:", re.I)
# An answer runs until the next answer marker, the end sentinel, the next
# numbered question heading, or a quoted-reply header block.
_Q_ANSWER_STOP_RE = re.compile(
    r"Answer\s*\[\s*iq-|\(End of questionnaire\)"
    r"|(?:^|\n)\s*(?:>+\s*)*\d+\.\s+[A-Z]{2}"
    r"|(?:^|\n)\s*(?:>+\s*)*From:\s",
)


def _questionnaire_subject(user_key: str) -> str:
    return f"{_Q_SUBJECT_PREFIX} — inbox-{user_key}"


def _load_questionnaire_parts(paths: dict) -> tuple[str, list[str], str]:
    """Split the questionnaire template into (intro, [questions], outro).
    Paragraphs starting '<n>.' are questions; text before the first is the
    intro, after the last the outro. A per-user copy in the share folder
    overrides the repo template, so a questionnaire can be customized."""
    from rocky import PROGRAM_DIR  # lazy
    src = paths["questionnaire"]
    if not src.exists():
        template = PROGRAM_DIR / "_templates" / "inbox_questionnaire.md"
        if template.exists():
            src.parent.mkdir(parents=True, exist_ok=True)
            src.write_text(template.read_text(encoding="utf-8"),
                           encoding="utf-8")
    if not src.exists():
        return "", [], ""
    text = src.read_text(encoding="utf-8")
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    intro_parts, questions, outro_parts = [], [], []
    for p in paras:
        if re.match(r"^\d+\.\s", p):
            questions.append(re.sub(r"^\d+\.\s*", "", p))
        elif not questions:
            intro_parts.append(p)
        else:
            outro_parts.append(p)
    return "\n\n".join(intro_parts), questions, "\n\n".join(outro_parts)


def _questionnaire_html(intro: str, questions: list[str], outro: str) -> str:
    """HTML body: intro, then each question with an answer box under it."""
    import html as _html

    def esc(s: str) -> str:
        return _html.escape(s).replace("\n", "<br>")

    font = "font-family:'Segoe UI',Arial,sans-serif;font-size:14px;"
    parts = [f"<div style=\"{font}color:#222;max-width:720px;\">",
             f"<p>{esc(intro)}</p>"]
    for i, q in enumerate(questions, start=1):
        m = re.match(r"^([A-Z][A-Z /&]+)\.\s*(.*)$", q, re.S)
        title, body = (m.group(1), m.group(2)) if m else ("", q)
        heading = f"{i}. {esc(title)}" if title else f"{i}."
        parts.append(
            f"<p style=\"margin:16px 0 4px 0;\"><b>{heading}</b> "
            f"{esc(body)}</p>"
            f"<div style=\"border:1px solid #b5b5b5;background:#fafafa;"
            f"padding:10px 12px;margin:4px 0 4px 0;min-height:52px;\">"
            f"<span style=\"color:#8a8a8a;\">Answer [iq-{i}]:</span>"
            f"<br><br><br></div>")
    # Sentinel goes BEFORE the outro: the reply parser reads each answer up
    # to the next stop marker, and the quoted-back outro would otherwise
    # bleed into the last answer (happened on Matt's reply, 2026-07-09).
    parts.append("<p style=\"color:#8a8a8a;\">(End of questionnaire)</p>")
    if outro:
        parts.append(f"<p style=\"margin-top:20px;\">{esc(outro)}</p>")
    parts.append("</div>")
    return "\n".join(parts)


def _clean_answer_text(text: str) -> str:
    """Strip quote prefixes / blank lines from an extracted answer chunk."""
    lines = []
    for ln in (text or "").splitlines():
        ln = re.sub(r"^\s*(?:>+\s*)+", "", ln).strip()
        if ln:
            lines.append(ln)
    return " ".join(lines).strip()


def parse_questionnaire_answers(body: str) -> dict[int, str]:
    """Extract {question_number: answer} from a reply body. Empty boxes
    (the unanswered original quoted back) parse to nothing."""
    answers: dict[int, str] = {}
    for m in _Q_ANSWER_MARKER_RE.finditer(body or ""):
        num = int(m.group(1))
        stop = _Q_ANSWER_STOP_RE.search(body, pos=m.end())
        chunk = body[m.end():stop.start()] if stop else body[m.end():]
        text = _clean_answer_text(chunk)
        if text and num not in answers:
            answers[num] = text[:2000]
    return answers


def questionnaire_cycle(config: dict, ucfg: dict, paths: dict,
                        resend: bool = False) -> dict:
    """Idempotent: not sent -> send it; sent but unanswered -> check
    rocky@'s inbox for the reply; answered -> quiet no-op."""
    from rocky import get_msal_app, acquire_token  # lazy
    from outbound import send_mail_guarded

    state = load_state(paths)
    rocky_email = config.get("rocky_email", "rocky@gallagherllp.com")
    user_addr = ucfg["mailbox"]
    subject = _questionnaire_subject(paths["user_key"])

    if state.get("q_answered_at") and not resend:
        log_event(paths, "questionnaire_noop", reason="already answered",
                  answered_at=state["q_answered_at"])
        return {"answered_at": state["q_answered_at"]}

    app = get_msal_app(config)
    token = acquire_token(app)

    # --- Send (once, or on --resend). --------------------------------------
    if not state.get("q_sent_at") or resend:
        intro, questions, outro = _load_questionnaire_parts(paths)
        if not questions:
            log_event(paths, "questionnaire_template_missing",
                      path=str(paths["questionnaire"]))
            return {"error": "questionnaire template missing/empty"}
        html = _questionnaire_html(intro, questions, outro)
        result = send_mail_guarded(token, rocky_email, [user_addr],
                                   subject, html, body_type="HTML")
        if not result.get("sent"):
            log_event(paths, "questionnaire_send_failed",
                      reason=result.get("reason"))
            return {"error": f"send failed: {result.get('reason')}"}
        state["q_sent_at"] = _now_iso()
        state.pop("q_answered_at", None)
        save_state(paths, state)
        log_comm(paths, "out", f"(questionnaire email, {len(questions)} "
                 f"questions, subject {subject!r})",
                 cohort_id="questionnaire")
        log_event(paths, "questionnaire_emailed", to=user_addr,
                  questions=len(questions))
        return {"sent": True, "questions": len(questions)}

    # --- Check for the reply in rocky@'s inbox. ----------------------------
    url = f"{GRAPH_API_BASE}/users/{rocky_email}/mailFolders/inbox/messages"
    params = {
        "$filter": f"receivedDateTime gt {state['q_sent_at']}",
        "$orderby": "receivedDateTime desc",
        "$top": "50",
        "$select": "id,subject,from,receivedDateTime,body,bodyPreview",
    }
    headers = {"Authorization": f"Bearer {token}",
               "Accept": "application/json",
               "Prefer": 'outlook.body-content-type="text"'}
    try:
        resp = requests.get(url, headers=headers, params=params, timeout=60)
    except requests.RequestException as e:
        log_event(paths, "questionnaire_check_failed", detail=str(e)[:200])
        return {"error": "network"}
    if resp.status_code != 200:
        log_event(paths, "questionnaire_check_failed",
                  detail=f"http {resp.status_code}: {resp.text[:200]}")
        return {"error": f"http {resp.status_code}"}

    reply = None
    for msg in resp.json().get("value", []):
        from_addr, _ = _addr_of(msg.get("from"))
        if from_addr == user_addr.lower() and \
                _Q_SUBJECT_PREFIX.lower() in (msg.get("subject") or "").lower():
            reply = msg
            break
    if reply is None:
        log_event(paths, "questionnaire_waiting", since=state["q_sent_at"])
        return {"waiting": True}

    body = ((reply.get("body") or {}).get("content")
            or reply.get("bodyPreview") or "")
    answers = parse_questionnaire_answers(body)
    log_comm(paths, "in", body[:6000], message_id=reply.get("id"),
             cohort_id="questionnaire", sender=user_addr)

    if not answers:
        log_event(paths, "questionnaire_reply_unparsed",
                  message_id=reply.get("id"),
                  note="reply found but no answers in boxes — James should "
                       "read it in communications.jsonl")
        return {"reply_found": True, "answers": 0}

    _, questions, _ = _load_questionnaire_parts(paths)
    lines = [f"# Questionnaire answers — {ucfg.get('display_name', paths['user_key'])}",
             f"\nReceived {reply.get('receivedDateTime')} from {user_addr}.\n"]
    for num in sorted(answers):
        qtext = questions[num - 1] if 0 < num <= len(questions) else f"(question {num})"
        lines.append(f"## {num}. {qtext}\n\n> {answers[num]}\n")
    paths["answers"].write_text("\n".join(lines) + "\n", encoding="utf-8")

    state["q_answered_at"] = _now_iso()
    save_state(paths, state)
    log_event(paths, "questionnaire_answered", answers=len(answers),
              of=len(questions), saved_to=str(paths["answers"]))

    ack = (f"Got it — thanks. I recorded {len(answers)} answer(s); they'll "
           f"shape the cleanup rules. Next you'll hear from me on Teams, "
           f"one batch at a time, and nothing moves until you approve it.")
    send_mail_guarded(token, rocky_email, [user_addr],
                      f"RE: {subject}", ack)
    log_comm(paths, "out", ack, cohort_id="questionnaire")
    return {"answers": len(answers)}


# =============================================================================
# Teams chat cycle
# =============================================================================

def _proposal_text(cohort: dict) -> str:
    header = f"[INBOX CLEANER — proposal {cohort['id']}]"
    body = f"{cohort.get('title')}\n\n{cohort.get('description')}"
    tops = cohort.get("top_senders")
    if tops:
        body += "\n\nTop senders: " + "; ".join(tops[:6])
    samples = cohort.get("sample_subjects")
    if samples:
        body += "\n\nMost common subjects: " + \
                "; ".join(f"“{s}”" for s in samples[:4])
    if cohort.get("target_folder") is None:
        ask = ("\n\nReply with the FOLDER NAME to file these under "
               "(I'll create it under your Inbox), or SKIP to leave them alone.")
    else:
        ask = (f"\n\nPlan: move them to \"{cohort['target_folder']}\". "
               f"Reply YES to approve or NO to skip. Nothing is ever deleted — "
               f"every move is logged and reversible.")
    return f"{header}\n\n{body}{ask}"


def _interpret_reply(text: str, needs_label: bool) -> tuple[str, str | None]:
    """Returns (decision, target_folder). decision in
    approved | declined | unclear."""
    t = (text or "").strip()
    low = t.lower().rstrip("!.")
    if low in _NO_WORDS:
        return "declined", None
    if needs_label:
        if not t:
            return "unclear", None
        if low in _YES_WORDS:
            return "unclear", None   # a bare YES doesn't name the folder
        return "approved", t[:120]
    if low in _YES_WORDS:
        return "approved", None
    return "unclear", None


def chat_cycle(config: dict, ucfg: dict, paths: dict) -> dict:
    """One Teams poll cycle. Logs every message both ways, resolves the
    pending proposal if the user replied, then proposes the next cohort.
    One live proposal at a time — a bare \"yes\" is never ambiguous.
    (The onboarding questionnaire is NOT sent here — it goes by email via
    questionnaire_cycle; Teams carries only the one-at-a-time proposals.)"""
    import teams
    from rocky import get_msal_app  # lazy — rocky imports us at dispatch time

    ensure_rules_file(paths, ucfg)
    state = load_state(paths)
    app = get_msal_app(config)
    try:
        token = teams.acquire_teams_token(app)
    except teams.TeamsNotEnabled as e:
        log_event(paths, "teams_unavailable", detail=str(e)[:300])
        log.error(f"{paths['tag']} Teams scopes not granted yet — see "
                  f"INBOX_CLEANER.md IT step 2. ({e})")
        return {"error": "teams_not_enabled"}

    rocky_upn = config.get("rocky_email", "rocky@gallagherllp.com")
    other_upn = (ucfg.get("chat_upn") or ucfg["mailbox"]).lower()
    # Observers (default: James) are IN the chat to watch and comment, but
    # only the mailbox owner's replies ever count as approvals.
    observers = ucfg.get("observers")
    if observers is None:
        observers = [config.get("user_email", "jbragdon@gallagherllp.com")]
    observers = [o.lower() for o in observers if o and o.lower() != other_upn]

    if not state.get("self_id"):
        state["self_id"] = teams.get_self_user_id(token)
    if not state.get("chat_id"):
        display = ucfg.get("display_name", paths["user_key"])
        if observers:
            topic = f"Rocky Inbox Cleaner — {display}"
            chat_id = teams.ensure_group_chat(
                token, topic, [rocky_upn, other_upn] + observers)
        else:
            chat_id = teams.ensure_one_on_one_chat(token, rocky_upn, other_upn)
        if not chat_id:
            log_event(paths, "chat_create_failed", other=other_upn)
            return {"error": "chat_create_failed"}
        state["chat_id"] = chat_id
        log_event(paths, "chat_created", chat_id=chat_id, with_user=other_upn,
                  observers=", ".join(observers) or "none")
        intro = (f"Hi {display.split()[0]} — this chat is where I'll propose "
                 f"inbox cleanup batches, one at a time. Reply YES to approve, "
                 f"NO to skip, or a folder name when I ask for one. Nothing "
                 f"is ever deleted, and nothing moves without your approval."
                 + (f" (Observers here can follow along, but only YOUR "
                    f"replies count as approvals.)" if observers else ""))
        sent = teams.send_chat_message(token, chat_id, intro)
        log_comm(paths, "out", intro, message_id=(sent or {}).get("id"))
    save_state(paths, state)
    chat_id = state["chat_id"]

    # Map member ids -> emails so we know WHO each reply is from. The owner
    # (the mailbox's user) is the only one whose replies decide cohorts.
    # Match tolerates the firm's legacy domain: the member's email may come
    # back as either @gallagherllp.com or @gejlaw.com.
    if not state.get("owner_id"):
        members = teams.list_chat_members(token, chat_id)
        state["members"] = {m["userId"]: m["email"]
                            for m in members if m.get("userId")}
        owner_addrs = {other_upn, teams.swap_legacy_domain(other_upn).lower()}
        state["owner_id"] = next(
            (m["userId"] for m in members if m.get("email") in owner_addrs),
            None)
        if state["owner_id"]:
            log_event(paths, "chat_owner_resolved", owner=other_upn)
        else:
            log_event(paths, "chat_owner_unresolved",
                      note="member list lacked the mailbox owner — "
                           "decisions deferred until resolved")
        save_state(paths, state)
    owner_id = state.get("owner_id")
    member_emails = state.get("members", {})

    # --- Ingest new replies (log EVERYTHING inbound). ----------------------
    msgs = teams.fetch_chat_messages(token, chat_id,
                                     since_iso=state.get("teams_cursor"))
    inbound: list[dict] = []
    for m in msgs:
        created = m.get("createdDateTime") or ""
        if created and created > (state.get("teams_cursor") or ""):
            state["teams_cursor"] = created
        sender = teams.message_sender_id(m)
        if sender and sender != state.get("self_id"):
            text = teams.message_text(m)
            log_comm(paths, "in", text, message_id=m.get("id"),
                     sender=member_emails.get(sender, sender))
            inbound.append({"text": text, "created": created,
                            "sender_id": sender})
    save_state(paths, state)

    cohorts = load_cohorts(paths)
    result: dict = {"inbound": len(inbound)}

    # --- Resolve the pending proposal, if any (strict protocol). -----------
    # ONLY the mailbox owner's replies decide cohorts — observers (James)
    # can chat freely without approving anything. If the owner id couldn't
    # be resolved, no reply decides anything (fail safe, logged above).
    pending = next((c for c in cohorts if c.get("status") == "proposed"), None)
    if pending and inbound:
        proposed_at = (pending.get("teams") or {}).get("proposed_at") or ""
        replies = [m for m in inbound if m["created"] > proposed_at
                   and owner_id is not None and m["sender_id"] == owner_id]
        if replies:
            reply = replies[0]
            needs_label = pending.get("target_folder") is None
            decision, label = _interpret_reply(reply["text"], needs_label)
            if decision == "approved":
                if label:
                    pending["target_folder"] = label
                pending["status"] = "approved"
                pending.setdefault("teams", {})["decided_at"] = _now_iso()
                pending["teams"]["reply"] = reply["text"][:300]
                record_decision_rule(paths, pending, "approved")
                ack = (f"Approved — {pending['count']:,} emails will move to "
                       f"\"{pending['target_folder']}\" on the next execution "
                       f"run. ({pending['id']})")
            elif decision == "declined":
                pending["status"] = "declined"
                pending.setdefault("teams", {})["decided_at"] = _now_iso()
                pending["teams"]["reply"] = reply["text"][:300]
                record_decision_rule(paths, pending, "declined")
                ack = f"Understood — I'll leave those alone. ({pending['id']})"
            else:
                pending["status"] = "needs_review"
                pending.setdefault("teams", {})["reply"] = reply["text"][:300]
                log_event(paths, "reply_unclear", cohort=pending["id"],
                          reply=reply["text"][:200])
                ack = (f"I couldn't read that as a yes/no"
                       f"{' or folder name' if needs_label else ''} for "
                       f"{pending['id']}, so I've set it aside for James to "
                       f"review. Nothing was moved.")
            sent = teams.send_chat_message(token, chat_id, ack)
            log_comm(paths, "out", ack, cohort_id=pending["id"],
                     message_id=(sent or {}).get("id"))
            log_event(paths, "cohort_decided", cohort=pending["id"],
                      status=pending["status"],
                      target=pending.get("target_folder"))
            save_cohorts(paths, cohorts)
            result["decided"] = {"id": pending["id"], "status": pending["status"]}
            pending = None

    # --- Propose the next cohort (one at a time). --------------------------
    still_pending = any(c.get("status") == "proposed" for c in cohorts)
    if not still_pending:
        # Most-obviously-non-work first (_PROPOSE_PRIORITY), biggest first
        # within a tier.
        nxt = min((c for c in cohorts if c.get("status") == "draft"),
                  key=lambda c: (_PROPOSE_PRIORITY.get(c.get("kind"), 6),
                                 -(c.get("count") or 0)),
                  default=None)
        if nxt is not None:
            text = _proposal_text(nxt)
            sent = teams.send_chat_message(token, chat_id, text)
            if sent:
                nxt["status"] = "proposed"
                nxt.setdefault("teams", {})["proposed_at"] = \
                    sent.get("createdDateTime") or _now_iso()
                nxt["teams"]["message_id"] = sent.get("id")
                save_cohorts(paths, cohorts)
                log_comm(paths, "out", text, cohort_id=nxt["id"],
                         message_id=sent.get("id"))
                log_event(paths, "cohort_proposed", cohort=nxt["id"],
                          count=nxt.get("count"))
                result["proposed"] = nxt["id"]

    return result


# =============================================================================
# Daily Claude rules update
# =============================================================================

def _read_jsonl_since(path: Path, since: datetime) -> list[dict]:
    if not path.exists():
        return []
    out = []
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            ts = rec.get("ts")
            try:
                if ts and datetime.fromisoformat(ts.replace("Z", "+00:00")) >= since:
                    out.append(rec)
            except ValueError:
                continue
    return out


def rules_update(config: dict, ucfg: dict, paths: dict, hours: int = 24) -> dict:
    """The one-Claude-call-a-day step: fold the day's decisions, chat
    traffic, and questionnaire answers into rules.md. Quiet day = no call."""
    ensure_rules_file(paths, ucfg)
    since = datetime.now(timezone.utc) - timedelta(hours=hours)
    events = _read_jsonl_since(paths["activity"], since)
    comms = _read_jsonl_since(paths["comms"], since)
    substantive = [e for e in events if e.get("event") in
                   ("cohort_decided", "rule_recorded", "questionnaire_emailed",
                    "questionnaire_answered", "reply_unclear")]
    if not substantive and not comms:
        log_event(paths, "rules_update_skipped", reason="no activity in window")
        return {"skipped": True}

    current = paths["rules"].read_text(encoding="utf-8")
    prompt = (
        f"You maintain the standing inbox-organization rules file for "
        f"{ucfg.get('display_name', paths['user_key'])} "
        f"(process inbox-{paths['user_key']}) at Gallagher LLP, a law firm. "
        f"Rocky (a virtual paralegal) proposes email cohorts over Teams; the "
        f"user approves or declines; approved/declined decisions become "
        f"durable rules for future automated inbox maintenance.\n\n"
        f"Here is the CURRENT rules file:\n\n---\n{current}\n---\n\n"
        f"Here are the last {hours}h of activity events (JSON lines):\n\n"
        + "\n".join(json.dumps(e, ensure_ascii=False) for e in substantive[-100:])
        + "\n\nAnd the last chat messages between Rocky and the user "
        f"(direction 'in' = the user speaking):\n\n"
        + "\n".join(json.dumps(c, ensure_ascii=False) for c in comms[-100:])
        + "\n\nRewrite the rules file so it fully reflects what was learned. "
        "Rules must stay conservative and auditable:\n"
        "- Keep every existing rule unless the day's activity clearly "
        "superseded it (then note the change inline, don't silently drop it).\n"
        "- Preserve the [C####] [approved/declined YYYY-MM-DD] provenance tags.\n"
        "- Extract any standing preferences the user expressed in chat or "
        "questionnaire answers into 'Standing preferences'.\n"
        "- Plain English, human-editable, keep the same three section headers.\n\n"
        "Return ONLY the complete updated markdown file, nothing else."
    )

    from anthropic import Anthropic
    client = Anthropic(api_key=config["anthropic_api_key"])
    try:
        resp = client.messages.create(
            model=CLAUDE_MODEL, max_tokens=RULES_MAX_TOKENS,
            messages=[{"role": "user", "content": prompt}])
        updated = resp.content[0].text.strip()
    except Exception as e:
        log_event(paths, "rules_update_failed", detail=str(e)[:300])
        return {"error": str(e)}

    if updated.startswith("```"):
        updated = re.sub(r"^```[a-z]*\n?|\n?```$", "", updated).strip()
    if "# Inbox rules" not in updated.splitlines()[0] and \
            "# Inbox rules" not in updated:
        log_event(paths, "rules_update_rejected",
                  reason="response missing header — keeping current file")
        return {"error": "malformed response"}

    paths["rules_history"].mkdir(parents=True, exist_ok=True)
    backup = paths["rules_history"] / \
        f"rules_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
    backup.write_text(current, encoding="utf-8")
    paths["rules"].write_text(updated + "\n", encoding="utf-8")
    log_event(paths, "rules_updated", backup=backup.name,
              events_considered=len(substantive), comms_considered=len(comms))
    result = {"updated": True, "backup": str(backup)}

    # Phase 2: tune matters.json from the day's chat (additive, guarded).
    # If anything changed, regenerate draft cohorts right away so the next
    # chat cycle proposes the updated batches — approval still gates moves.
    tune = _tune_matters(config, ucfg, paths, comms)
    if tune.get("applied"):
        result["matters_ops"] = tune["applied"]
        if paths["snapshot"].exists():
            result["reanalyzed"] = analyze(config, ucfg, paths)
    return result


# =============================================================================
# Nightly matters.json tuning — Rocky refines the deal list from the chat
# =============================================================================
# Runs inside --rules-update. A second Claude call reads the day's chat and
# matter-cohort decisions and proposes STRUCTURED edits to matters.json:
# add_matter / update_keywords / set_closed — never delete. The code
# validates and applies them deterministically, backs up the old file, logs
# every change (the daily digest reports it), and re-runs analyze so the
# next chat cycle proposes the updated batches. Safe by construction: a
# changed rule only produces a fresh DRAFT cohort — nothing moves until the
# user approves it on Teams.

def _apply_matter_ops(data: dict, ops: list[dict]) -> list[str]:
    """Validate + apply ops to the matters structure in place. Returns
    plain-English descriptions of what was actually applied (invalid or
    unresolvable ops are skipped silently — conservative)."""
    applied: list[str] = []
    clients = data.setdefault("clients", [])

    def find(client: str, matter: str):
        for cl in clients:
            if (cl.get("client") or "").strip().lower() == client.lower():
                for m in cl.get("matters") or []:
                    if (m.get("matter") or "").strip().lower() == matter.lower():
                        return cl, m
                return cl, None
        return None, None

    for op in ops or []:
        kind = op.get("op")
        client = (op.get("client") or "").strip()
        matter = (op.get("matter") or "").strip()
        if not (client and matter):
            continue
        cl, m = find(client, matter)
        kws = [k.strip() for k in (op.get("keywords") or [])
               if isinstance(k, str) and k.strip()]
        if kind == "add_matter":
            if m is not None or not kws:
                continue
            if cl is None:
                cl = {"client": client, "matters": []}
                clients.append(cl)
            cl.setdefault("matters", []).append(
                {"matter": matter, "keywords": kws,
                 "closed": bool(op.get("closed"))})
            applied.append(f"added {client} — {matter} "
                           f"(keywords: {', '.join(kws)})")
        elif kind == "update_keywords":
            if m is None or not kws:
                continue
            old = list(m.get("keywords") or [])
            if kws == old:
                continue
            m["keywords"] = kws
            applied.append(f"keywords for {client} — {matter}: "
                           f"{old} → {kws}")
        elif kind == "set_closed":
            if m is None or "closed" not in op:
                continue
            want = bool(op["closed"])
            if bool(m.get("closed")) == want:
                continue
            m["closed"] = want
            applied.append(f"{client} — {matter} marked "
                           f"{'closed' if want else 'open'}")
    return applied


def _tune_matters(config: dict, ucfg: dict, paths: dict,
                  comms: list[dict]) -> dict:
    """One guarded Claude call: propose matters.json edits grounded in the
    day's chat. No file = no tuning; no clear signal = no edits."""
    p = paths.get("matters")
    if not p or not p.exists() or not comms:
        return {}
    noun = _matter_noun(paths, ucfg)
    name = ucfg.get("display_name", paths["user_key"])
    current = p.read_text(encoding="utf-8")
    matter_cohorts = [
        {"title": c.get("title"), "status": c.get("status"),
         "count": c.get("count"),
         "keywords": (c.get("match") or {}).get("matter_keywords")}
        for c in load_cohorts(paths) if c.get("kind") == "matter"]
    chat = [{"direction": c.get("direction"), "sender": c.get("sender"),
             "text": (c.get("text") or "")[:400]} for c in comms[-40:]]

    prompt = (
        f"You maintain the {noun} list that drives automated inbox filing "
        f"for {name} at Gallagher LLP. The file below defines their "
        f"{noun}s; message subjects are matched against each {noun}'s "
        f"keywords (whole-word, case-insensitive) and matching "
        f"conversations are filed to that {noun}'s folder — but only after "
        f"{name} approves each batch on Teams, so your edits can never "
        f"move mail directly.\n\n"
        f"Current matters.json:\n---\n{current}\n---\n\n"
        f"Current {noun} batches and their status (JSON):\n"
        f"{json.dumps(matter_cohorts, ensure_ascii=False)}\n\n"
        f"Today's chat between Rocky and {name} "
        f"(direction 'in' = {name} speaking):\n"
        f"{json.dumps(chat, ensure_ascii=False)}\n\n"
        "Propose edits ONLY where the chat clearly supports them:\n"
        f"- {name} names a NEW {noun} not in the list → add_matter, with "
        "distinctive subject keywords taken from their words.\n"
        f"- {name} says a batch caught the wrong mail or its keywords "
        "need adjusting → update_keywords (the full replacement list).\n"
        f"- {name} says a {noun} closed (or reopened) → set_closed.\n"
        "Rules: never remove a matter; never invent names or keywords not "
        "grounded in the chat; prefer distinctive multi-word phrases over "
        "common English words; when in doubt, propose nothing.\n\n"
        'Return ONLY JSON, no prose: {"ops": [...]} where each op is\n'
        '{"op": "add_matter", "client": "...", "matter": "...", '
        '"keywords": ["..."], "closed": false}\n'
        '{"op": "update_keywords", "client": "...", "matter": "...", '
        '"keywords": ["..."]}\n'
        '{"op": "set_closed", "client": "...", "matter": "...", '
        '"closed": true}\n'
        'No changes = {"ops": []}.'
    )
    try:
        from anthropic import Anthropic
        client = Anthropic(api_key=config["anthropic_api_key"])
        resp = client.messages.create(
            model=CLAUDE_MODEL, max_tokens=2000,
            messages=[{"role": "user", "content": prompt}])
        raw = resp.content[0].text.strip()
        if raw.startswith("```"):
            raw = re.sub(r"^```[a-z]*\n?|\n?```$", "", raw).strip()
        ops = (json.loads(raw).get("ops") or [])
    except Exception as e:
        log_event(paths, "matters_tune_failed", detail=str(e)[:300])
        return {}
    if not ops:
        return {}

    try:
        data = json.loads(current)
    except json.JSONDecodeError as e:
        log_event(paths, "matters_tune_failed",
                  detail=f"matters.json unparseable: {e}")
        return {}
    applied = _apply_matter_ops(data, ops)
    if not applied:
        return {}

    paths["rules_history"].mkdir(parents=True, exist_ok=True)
    backup = paths["rules_history"] / \
        f"matters_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    backup.write_text(current, encoding="utf-8")
    p.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n",
                 encoding="utf-8")
    log_event(paths, "matters_updated", ops=applied, backup=backup.name)
    return {"applied": applied, "backup": str(backup)}


# =============================================================================
# Daily digest — plain-English activity summary, emailed to owner + observers
# =============================================================================
# Deterministic facts come first (counts from the window's activity, comms,
# and move logs); one Claude call phrases them plainly. If the API call
# fails, a deterministic summary goes out instead — the digest never blocks
# on Claude. A quiet window sends nothing. Recipients: the mailbox owner +
# observers (default James), the same audience as the approval chat.

# Bookkeeping events that don't count as "activity" for digest purposes —
# in particular the digest's own events, so a sent digest never makes the
# next day look active.
_DIGEST_IGNORED_EVENTS = {
    "digest_sent", "digest_skipped", "digest_send_failed",
    "questionnaire_noop", "questionnaire_waiting",
    "rules_update_skipped", "execute_skipped", "analyze_skipped",
}

_DIGEST_STATUS_LABELS = {
    "draft": "drafted, not yet proposed",
    "proposed": "proposed on Teams, awaiting a reply",
    "approved": "approved, queued to move",
    "declined": "declined — left alone",
    "needs_review": "set aside for James to review",
    "executed": "moved",
}


def _digest_facts(paths: dict, since: datetime) -> dict:
    """Aggregate the window's logs into everything the digest reports."""
    events = [e for e in _read_jsonl_since(paths["activity"], since)
              if e.get("event") not in _DIGEST_IGNORED_EVENTS]
    comms = _read_jsonl_since(paths["comms"], since)
    moves = _read_jsonl_since(paths["moves"], since)

    cohorts = load_cohorts(paths)
    by_id = {c.get("id"): c for c in cohorts}

    def _title(cid: str | None) -> str:
        c = by_id.get(cid) or {}
        return c.get("title") or (cid or "?")

    moved_by_cohort: dict[str, int] = {}
    for m in moves:
        t = _title(m.get("cohort"))
        moved_by_cohort[t] = moved_by_cohort.get(t, 0) + 1

    by_event: dict[str, int] = {}
    for e in events:
        by_event[e.get("event", "?")] = by_event.get(e.get("event", "?"), 0) + 1

    by_status: dict[str, int] = {}
    for c in cohorts:
        s = c.get("status", "?")
        by_status[s] = by_status.get(s, 0) + 1

    pending = next((c for c in cohorts if c.get("status") == "proposed"), None)
    return {
        "events": events,
        "comms": comms,
        "by_event": by_event,
        "by_status": by_status,
        "moved_total": len(moves),
        "moved_by_cohort": moved_by_cohort,
        "move_failures": by_event.get("move_failed", 0),
        "decided": [{"batch": _title(e.get("cohort")), "status": e.get("status")}
                    for e in events if e.get("event") == "cohort_decided"],
        "proposed": [{"batch": _title(e.get("cohort")), "count": e.get("count")}
                     for e in events if e.get("event") == "cohort_proposed"],
        "pending": ({"title": pending.get("title"),
                     "count": pending.get("count"),
                     "target": pending.get("target_folder")}
                    if pending else None),
        "matters_changes": [op for e in events
                            if e.get("event") == "matters_updated"
                            for op in (e.get("ops") or [])],
    }


def _digest_fallback_md(facts: dict) -> str:
    """Plain deterministic summary used when the Claude call fails."""
    lines = ["**What happened**"]
    if facts["moved_total"]:
        lines.append(f"- {facts['moved_total']:,} email(s) were filed into "
                     f"folders:")
        for title, n in sorted(facts["moved_by_cohort"].items(),
                               key=lambda kv: -kv[1]):
            lines.append(f"- {n:,} → {title}")
    for d in facts["decided"]:
        label = _DIGEST_STATUS_LABELS.get(d["status"], d["status"])
        lines.append(f"- Batch \"{d['batch']}\": {label}.")
    for p in facts["proposed"]:
        n = p.get("count")
        lines.append(f"- New batch proposed on Teams: \"{p['batch']}\""
                     + (f" ({n:,} emails)." if isinstance(n, int) else "."))
    if facts["by_event"].get("questionnaire_answered"):
        lines.append("- Your questionnaire answers were received and saved.")
    if facts["by_event"].get("rules_updated"):
        lines.append("- The standing rules file was updated.")
    for ch in facts.get("matters_changes") or []:
        lines.append(f"- Deal/matter list updated: {ch}. A new or revised "
                     f"batch proposal will follow.")
    if facts["by_event"].get("snapshot_done"):
        lines.append("- A read-only snapshot of the mailbox was taken.")
    if facts["move_failures"]:
        lines.append(f"- {facts['move_failures']:,} move(s) failed and will "
                     f"be retried; nothing was lost.")
    if len(lines) == 1:
        lines.append("- Housekeeping only — nothing that needs your "
                     "attention.")
    if facts["pending"]:
        p = facts["pending"]
        n = p.get("count")
        lines += ["", "**Waiting on you**",
                  f"- \"{p['title']}\""
                  + (f" ({n:,} emails)" if isinstance(n, int) else "")
                  + " is waiting for your YES/NO on Teams."]
    return "\n".join(lines)


def _digest_summary_md(config: dict, ucfg: dict, facts: dict,
                       hours: int, noun: str = "case") -> tuple[str, bool]:
    """One Claude call -> plain-English markdown. Returns (md, used_claude)."""
    name = ucfg.get("display_name", "the user")
    slim = {k: facts[k] for k in ("by_event", "by_status", "moved_total",
                                  "moved_by_cohort", "move_failures",
                                  "decided", "proposed", "pending",
                                  "matters_changes")}
    comms = [{"direction": c.get("direction"), "sender": c.get("sender"),
              "text": (c.get("text") or "")[:300]}
             for c in facts["comms"][-20:]]
    prompt = (
        f"You write the daily digest of Rocky's Inbox Cleaner — an email "
        f"cleanup assistant at Gallagher LLP, a law firm — for "
        f"{name}, whose mailbox it is. Address {name} as 'you'. James "
        f"Bragdon (the attorney supervising the process) receives the same "
        f"email. {name}'s practice is organized around {noun}s — talk "
        f"about {noun}s and projects, never 'cases' or 'matters' unless "
        f"that IS the noun.\n\n"
        f"Aggregated activity for the last {hours} hours (JSON):\n"
        f"{json.dumps(slim, ensure_ascii=False)}\n\n"
        f"Recent messages between Rocky and {name} (JSON):\n"
        f"{json.dumps(comms, ensure_ascii=False)}\n\n"
        "Write a short plain-English summary. Rules:\n"
        "- Say 'batch', never 'cohort'; no jargon, no JSON keys, no event "
        "names — describe what happened the way a person would.\n"
        "- Report ONLY facts present in the data above; omit anything "
        "ambiguous. Never invent numbers.\n"
        "- Formatting: only paragraphs, '- ' bullet lines, and bold "
        "'**Heading**' lines on their own (the renderer supports nothing "
        "else — no #, no links, no tables).\n"
        "- Sections: '**What happened**' (always); '**Waiting on you**' "
        "only if a batch is awaiting a reply; '**By the numbers**' only "
        "if emails were actually moved.\n"
        f"- If matters_changes is non-empty, say plainly that the {noun} "
        "list was updated and how (it means a new or revised batch "
        "proposal will follow for approval).\n"
        "- Under 200 words. Return ONLY the markdown, nothing else.\n"
    )
    try:
        from anthropic import Anthropic
        client = Anthropic(api_key=config["anthropic_api_key"])
        resp = client.messages.create(
            model=CLAUDE_MODEL, max_tokens=1500,
            messages=[{"role": "user", "content": prompt}])
        md = resp.content[0].text.strip()
        if md.startswith("```"):
            md = re.sub(r"^```[a-z]*\n?|\n?```$", "", md).strip()
        if md:
            return md, True
    except Exception as e:
        log.warning(f"[inbox-cleaner] digest Claude call failed, using "
                    f"deterministic summary: {str(e)[:200]}")
    return _digest_fallback_md(facts), False


def _digest_html(ucfg: dict, body_md: str, hours: int, facts: dict) -> str:
    """The digest email, styled like Rocky's other digests."""
    from rocky import _md_section_to_html  # lazy
    today = datetime.now(timezone.utc).strftime("%B %d, %Y")
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    name = ucfg.get("display_name", "")
    body_html = _md_section_to_html(body_md)

    stats = [f"Window: last {hours} hours"]
    if facts["moved_total"]:
        stats.append(f"<strong>{facts['moved_total']:,}</strong> email(s) filed")
    if facts["pending"]:
        stats.append("<strong>1</strong> batch awaiting your reply")
    stats_str = " &middot; ".join(stats)

    return f"""<!DOCTYPE html>
<html lang="en" xmlns="http://www.w3.org/1999/xhtml">
<head>
    <meta charset="utf-8"/>
    <meta name="viewport" content="width=device-width, initial-scale=1.0"/>
    <title>Rocky Inbox Cleaner Digest — {today}</title>
    <!--[if mso]>
    <style type="text/css">
        table {{border-collapse:collapse;}}
        td {{font-family:Segoe UI,Arial,sans-serif;}}
    </style>
    <![endif]-->
</head>
<body style="margin:0;padding:0;background-color:#f7f8fa;font-family:Segoe UI,Calibri,Arial,sans-serif;">
    <table width="100%" cellpadding="0" cellspacing="0" border="0"
           style="background-color:#f7f8fa;">
        <tr><td align="center" style="padding:24px 16px;">
            <table width="640" cellpadding="0" cellspacing="0" border="0"
                   style="background-color:#ffffff;border-radius:8px;
                          box-shadow:0 1px 3px rgba(0,0,0,0.08);max-width:640px;">
                <!-- Header -->
                <tr><td style="background-color:#1a202c;padding:28px 32px;
                               border-radius:8px 8px 0 0;">
                    <table width="100%" cellpadding="0" cellspacing="0" border="0">
                        <tr>
                            <td width="64" valign="top" style="padding-right:16px;">
                                <img src="cid:rocky_icon" width="56" height="56"
                                     alt="Rocky"
                                     style="display:block;border-radius:8px;"/>
                            </td>
                            <td valign="middle">
                                <h1 style="margin:0;font-size:22px;font-weight:700;
                                           color:#ffffff;line-height:1.2;">
                                    Rocky's Inbox Cleaner Digest</h1>
                                <p style="margin:4px 0 0 0;font-size:15px;
                                          color:#a0aec0;font-weight:500;">
                                    {name} &middot; {today}</p>
                                <p style="margin:4px 0 0 0;font-size:12px;
                                          color:#718096;font-style:italic;">
                                    What your inbox cleanup did today.</p>
                            </td>
                        </tr>
                    </table>
                </td></tr>
                <!-- Summary bar -->
                <tr><td style="background-color:#edf2f7;padding:12px 32px;
                               border-bottom:1px solid #e2e8f0;">
                    <p style="margin:0;font-size:13px;color:#4a5568;">
                        Generated {now_str} &middot; {stats_str}</p>
                </td></tr>
                <!-- Body -->
                <tr><td style="padding:20px 32px 24px 32px;">
                    {body_html}
                </td></tr>
                <!-- Footer -->
                <tr><td style="background-color:#f7f8fa;padding:16px 32px;
                               border-top:1px solid #e2e8f0;
                               border-radius:0 0 8px 8px;">
                    <p style="margin:0;font-size:11px;color:#a0aec0;text-align:center;">
                        Generated automatically by Rocky. Nothing is ever
                        deleted — newsletters go to Deleted Items (recoverable)
                        and everything else moves to named folders; every move
                        is logged and reversible. Nothing moves without your
                        approval on Teams.</p>
                </td></tr>
            </table>
        </td></tr>
    </table>
</body>
</html>"""


def digest_cycle(config: dict, ucfg: dict, paths: dict, hours: int = 24) -> dict:
    """Build and email the plain-English daily digest. Quiet window = no
    email (and no Claude call)."""
    from rocky import get_msal_app, acquire_token, ROCKY_ICON_PATH  # lazy
    from outbound import send_mail_guarded

    since = datetime.now(timezone.utc) - timedelta(hours=hours)
    facts = _digest_facts(paths, since)
    if not facts["events"] and not facts["comms"] and not facts["moved_total"]:
        log_event(paths, "digest_skipped", reason="no activity in window")
        return {"skipped": True}

    body_md, used_claude = _digest_summary_md(
        config, ucfg, facts, hours, noun=_matter_noun(paths, ucfg))
    html = _digest_html(ucfg, body_md, hours, facts)

    recipients = [ucfg["mailbox"]]
    observers = ucfg.get("observers")
    if observers is None:
        observers = [config.get("user_email", "jbragdon@gallagherllp.com")]
    for obs in observers:
        if obs and obs.lower() not in (r.lower() for r in recipients):
            recipients.append(obs)

    icon_attachment = []
    if ROCKY_ICON_PATH.exists():
        icon_attachment = [{"path": str(ROCKY_ICON_PATH),
                            "name": "rocky_icon.png",
                            "contentId": "rocky_icon"}]

    app = get_msal_app(config)
    token = acquire_token(app)
    rocky_email = config.get("rocky_email", "rocky@gallagherllp.com")
    today = datetime.now(timezone.utc).strftime("%B %d, %Y")
    subject = (f"Rocky Inbox Cleaner Digest — "
               f"{ucfg.get('display_name', paths['user_key'])} — {today}")
    result = send_mail_guarded(token, rocky_email, recipients, subject, html,
                               body_type="HTML", attachments=icon_attachment)
    if not result.get("sent"):
        log_event(paths, "digest_send_failed", reason=result.get("reason"))
        return {"error": f"send failed: {result.get('reason')}"}

    log_comm(paths, "out", body_md, cohort_id="digest")
    log_event(paths, "digest_sent", to=recipients, hours=hours,
              used_claude=used_claude, moved=facts["moved_total"])
    return {"sent": True, "to": recipients, "used_claude": used_claude}


# =============================================================================
# Internal-sweep report — what's inside, emailed BEFORE anything is proposed
# =============================================================================

def internal_report(config: dict, ucfg: dict, paths: dict) -> dict:
    """Email the mailbox owner (+ observers) a detailed breakdown of the
    internal-sweep subgroups before they hit the Teams chat. Built after a
    single 51k 'internal' batch proved unreviewable (Matt, 2026-07-12).
    Deterministic — no Claude call; must run where the snapshot lives."""
    from rocky import (get_msal_app, acquire_token, ROCKY_ICON_PATH,
                       _md_section_to_html)  # lazy
    from outbound import send_mail_guarded

    records = load_snapshot(paths)
    if not records:
        log_event(paths, "internal_report_skipped", reason="no snapshot")
        return {"error": "no snapshot — run --snapshot first"}
    senders, convs = build_ledgers(records)
    drafts = propose_cohorts(records, senders, convs, ucfg, paths)
    groups = [d for d in drafts if d["kind"].startswith("internal_")]
    if not groups:
        log_event(paths, "internal_report_skipped", reason="no internal groups")
        return {"error": "no internal subgroups drafted"}

    noun = _matter_noun(paths, ucfg)
    display = ucfg.get("display_name", paths["user_key"])
    total = sum(d["count"] for d in groups)

    md = [
        (f"Before anything is approved, here's exactly what's inside the "
         f"internal-mail cleanup: {total:,} emails, split into the "
         f"{len(groups)} groups below. I'll propose each group separately "
         f"in our Teams chat — most obviously non-work mail first — and "
         f"nothing moves without your YES on that specific group."),
        "",
        (f"If anything below is really {noun} mail — for example internal "
         f"emails about a live {noun} that never name it in the subject — "
         f"reply in the Teams chat and I'll re-route those to the "
         f"{noun}'s folder instead of archiving them."),
        "",
    ]
    for i, d in enumerate(groups, 1):
        md.append(f"**{i}. {d['title']} — {d['count']:,} emails**")
        md.append(d["description"])
        if d.get("top_senders"):
            md.append("- Top senders: " + "; ".join(d["top_senders"][:8]))
        if d.get("sample_subjects"):
            md.append("- Common subjects: "
                      + "; ".join(f"“{s}”"
                                  for s in d["sample_subjects"][:5]))
        md.append("")
    md.append("Everything lands in named, searchable folders — nothing is "
              "ever deleted, and every move is logged and reversible.")
    body_md = "\n".join(md)

    report_path = paths["share"] / "internal_report.md"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(body_md + "\n", encoding="utf-8")

    today = datetime.now(timezone.utc).strftime("%B %d, %Y")
    body_html = _md_section_to_html(body_md)
    html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"/></head>
<body style="margin:0;padding:0;background-color:#f7f8fa;font-family:Segoe UI,Calibri,Arial,sans-serif;">
<table width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#f7f8fa;">
<tr><td align="center" style="padding:24px 16px;">
<table width="640" cellpadding="0" cellspacing="0" border="0"
       style="background-color:#ffffff;border-radius:8px;box-shadow:0 1px 3px rgba(0,0,0,0.08);max-width:640px;">
<tr><td style="background-color:#1a202c;padding:28px 32px;border-radius:8px 8px 0 0;">
    <table width="100%" cellpadding="0" cellspacing="0" border="0"><tr>
        <td width="64" valign="top" style="padding-right:16px;">
            <img src="cid:rocky_icon" width="56" height="56" alt="Rocky"
                 style="display:block;border-radius:8px;"/></td>
        <td valign="middle">
            <h1 style="margin:0;font-size:22px;font-weight:700;color:#ffffff;line-height:1.2;">
                What's inside the internal-mail cleanup</h1>
            <p style="margin:4px 0 0 0;font-size:15px;color:#a0aec0;font-weight:500;">
                {display} &middot; {today}</p>
        </td></tr></table>
</td></tr>
<tr><td style="padding:20px 32px 24px 32px;">{body_html}</td></tr>
<tr><td style="background-color:#f7f8fa;padding:16px 32px;border-top:1px solid #e2e8f0;border-radius:0 0 8px 8px;">
    <p style="margin:0;font-size:11px;color:#a0aec0;text-align:center;">
        Generated automatically by Rocky. Nothing moves without your
        approval on Teams; nothing is ever deleted.</p>
</td></tr>
</table></td></tr></table></body></html>"""

    recipients = [ucfg["mailbox"]]
    observers = ucfg.get("observers")
    if observers is None:
        observers = [config.get("user_email", "jbragdon@gallagherllp.com")]
    for obs in observers:
        if obs and obs.lower() not in (r.lower() for r in recipients):
            recipients.append(obs)
    icon = []
    if ROCKY_ICON_PATH.exists():
        icon = [{"path": str(ROCKY_ICON_PATH), "name": "rocky_icon.png",
                 "contentId": "rocky_icon"}]
    app = get_msal_app(config)
    token = acquire_token(app)
    rocky_email = config.get("rocky_email", "rocky@gallagherllp.com")
    subject = (f"Rocky Inbox Cleaner — what's inside the internal-mail "
               f"cleanup ({total:,} emails, {len(groups)} groups)")
    result = send_mail_guarded(token, rocky_email, recipients, subject, html,
                               body_type="HTML", attachments=icon)
    if not result.get("sent"):
        log_event(paths, "internal_report_send_failed",
                  reason=result.get("reason"))
        return {"error": f"send failed: {result.get('reason')}"}
    log_comm(paths, "out", body_md[:6000], cohort_id="internal_report")
    log_event(paths, "internal_report_sent", to=recipients, groups=len(groups),
              total=total, saved_to=str(report_path))
    return {"sent": True, "groups": len(groups), "total": total}


# =============================================================================
# Stage 4 — execute (guarded)
# =============================================================================
# Two write paths, per-user (config inbox_users.<key>.write_via):
#   "app" (default)  — application token; needs Mail.ReadWrite (Application)
#                      + admin consent, fenced by the Application Access
#                      Policy.
#   "delegated"      — rocky@'s own delegated token; needs only an Exchange
#                      Full Access delegation on the target mailbox (the
#                      same mechanism as James's mailbox) — no Azure change.
#                      Chosen for Matt 2026-07-09 as the lighter IT ask.

# Kept separate from rocky.GRAPH_SCOPES so mail commands never prompt for
# write consent. Verified consented in rocky@'s grant 2026-07-05.
DELEGATED_WRITE_SCOPES = ["Mail.ReadWrite.Shared"]


def acquire_delegated_write_token(config: dict) -> str:
    """Delegated token carrying Mail.ReadWrite.Shared, acquired silently
    from rocky@'s cached sign-in (same MSAL cache as everything else)."""
    from rocky import get_msal_app  # lazy
    app = get_msal_app(config)
    accounts = app.get_accounts()
    result = app.acquire_token_silent(
        DELEGATED_WRITE_SCOPES, account=accounts[0]) if accounts else None
    if hasattr(app, "_save_cache"):
        app._save_cache()
    if not result or "access_token" not in result:
        log.error(
            "[inbox-cleaner] could not acquire a delegated write token "
            "(Mail.ReadWrite.Shared) from rocky@'s cached sign-in. Run any "
            "delegated command (e.g. --questionnaire) to refresh the "
            "sign-in, then retry.")
        sys.exit(1)
    return result["access_token"]


def _write_token(config: dict, ucfg: dict) -> str:
    """The token for execute-phase Graph calls, per the user's write_via."""
    if (ucfg.get("write_via") or "app").strip().lower() == "delegated":
        return acquire_delegated_write_token(config)
    from rocky import acquire_app_token  # lazy
    return acquire_app_token(config)

def _cohort_message_ids(cohort: dict, records: list[dict]) -> list[str]:
    """Message ids matching a cohort's criteria + age floor, from the snapshot.

    Match criteria are stored as RULES (keywords, senders, domains), not
    resolved message ids, so they stay stable across snapshots; this
    function re-resolves them against the current snapshot."""
    match = cohort.get("match", {})
    floor = int(cohort.get("age_floor_days") or 0)
    cutoff = (datetime.now(timezone.utc) - timedelta(days=floor)) \
        .strftime("%Y-%m-%dT%H:%M:%SZ") if floor else None
    senders = set(match.get("senders") or [])
    domains = set(match.get("domains") or [])
    conv_keys = set(match.get("conversation_keys") or [])
    purely_internal = bool(match.get("purely_internal"))
    # Matter cohorts match whole conversations whose subjects name the
    # matter — resolve keywords to conversation keys against this snapshot
    # (honoring the participant-corroboration guard when the rule has one).
    if match.get("matter_keywords"):
        conv_keys |= _matter_conv_keys(records, match["matter_keywords"],
                                       match.get("participant_domains"))
        if not conv_keys:
            return []

    ids = []
    for r in records:
        if cutoff and (r.get("received") or "") >= cutoff:
            continue
        key = _conv_key_of(r)
        if conv_keys:
            if key not in conv_keys:
                continue
        elif senders or domains:
            frm = r.get("from") or ""
            if frm not in senders and not (domains and _domain_in(frm, domains)):
                continue
            if purely_internal and not all(
                    _is_internal(a) for a in
                    ([frm] + (r.get("to") or []) + (r.get("cc") or []))
                    if a):
                continue
            # Recipient-count bounds split the internal sweep's broadcast
            # vs person-to-person subgroups.
            nrec = len(r.get("to") or []) + len(r.get("cc") or [])
            if match.get("min_recipients") is not None and \
                    nrec < int(match["min_recipients"]):
                continue
            if match.get("max_recipients") is not None and \
                    nrec > int(match["max_recipients"]):
                continue
        else:
            continue
        ids.append(r["id"])
    return ids


def _ensure_folder(token: str, mailbox: str, path_str: str, live: bool,
                   paths: dict) -> str | None:
    """Resolve 'A\\B' under the Inbox, creating missing segments (live only).
    Returns the folder id, or None (dry-run for a not-yet-existing folder)."""
    parent = "inbox"
    for seg in [s for s in path_str.replace("/", "\\").split("\\") if s.strip()]:
        url = f"{GRAPH_API_BASE}/users/{mailbox}/mailFolders/{parent}/childFolders"
        resp = _graph_get(url, token, {"$top": "200", "$select": "id,displayName"})
        found = None
        if resp is not None and resp.status_code == 200:
            for f in resp.json().get("value", []):
                if (f.get("displayName") or "").lower() == seg.strip().lower():
                    found = f["id"]
                    break
        if not found:
            if not live:
                return None
            create = _graph_post(url, token, {"displayName": seg.strip()})
            if create is None or create.status_code not in (200, 201):
                log_event(paths, "folder_create_failed", folder=seg,
                          detail=(create.text[:200] if create is not None else "none"))
                return None
            found = create.json()["id"]
            log_event(paths, "folder_created", folder=seg)
        parent = found
    return parent


def _load_moved_ids(paths: dict) -> set[str]:
    moved = set()
    if paths["moves"].exists():
        with open(paths["moves"], "r", encoding="utf-8") as f:
            for line in f:
                try:
                    moved.add(json.loads(line)["id"])
                except (json.JSONDecodeError, KeyError):
                    continue
    return moved


def execute(token: str, config: dict, ucfg: dict, paths: dict,
            live: bool = False, limit: int | None = None) -> dict:
    """Execute approved cohorts. Dry-run by default. Live mode needs the
    Mail.ReadWrite application permission (IT gate) — until then every move
    call would 403, so don't pass --live before that's granted. Moves are
    checkpointed to moves.jsonl (with source folder) for undo-by-replay."""
    mailbox = ucfg["mailbox"]
    # Exclusions also hold at execute time, so a cohort approved BEFORE an
    # exclusion was learned can't take that sender's mail afterwards.
    records = _drop_excluded(load_snapshot(paths), *load_chat_exclusions(paths))
    cohorts = load_cohorts(paths)
    approved = [c for c in cohorts if c.get("status") == "approved"]
    if not approved:
        log_event(paths, "execute_skipped", reason="no approved cohorts")
        return {"moved": 0, "planned": 0}

    moved_ids = _load_moved_ids(paths)
    log_event(paths, "execute_start", live=live, cohorts=len(approved),
              already_moved=len(moved_ids), limit=limit)

    # Claim messages in kind-priority order across ALL cohorts (matters
    # first, internal sweep last) so a message matching two cohorts goes
    # where the higher-priority rule says — an approved internal sweep can
    # never take mail a matter cohort matches. Declined cohorts claim too
    # ("I'll leave those alone") but are never executed.
    def _prio(c: dict):
        return (_KIND_PRIORITY.get(c.get("kind"), 9),
                c.get("seq", 10 ** 9), -(c.get("count") or 0))

    claims: dict[str, list[str]] = {}
    claimed: set[str] = set()
    for c in sorted(cohorts, key=_prio):
        c_ids = [i for i in _cohort_message_ids(c, records)
                 if i not in moved_ids and i not in claimed]
        claims[c["id"]] = c_ids
        claimed.update(c_ids)

    total_planned = 0
    total_moved = 0
    token_refreshed = time.monotonic()
    for cohort in sorted(approved, key=_prio):
        ids = claims.get(cohort["id"], [])
        total_planned += len(ids)
        if cohort.get("disposition") == "deleteditems":
            dest: str | None = "deleteditems"
        else:
            dest = _ensure_folder(token, mailbox, cohort["target_folder"],
                                  live, paths)
        log_event(paths, "cohort_plan", cohort=cohort["id"],
                  to_move=len(ids), target=cohort.get("target_folder"),
                  dest_resolved=bool(dest), live=live)
        if not live:
            continue
        if not dest:
            log_event(paths, "cohort_execute_skipped", cohort=cohort["id"],
                      reason="destination folder unresolved")
            continue

        done = 0
        for mid in ids:
            if limit is not None and total_moved >= limit:
                break
            resp = _graph_post(
                f"{GRAPH_API_BASE}/users/{mailbox}/messages/{mid}/move",
                token, {"destinationId": dest})
            if resp is not None and resp.status_code == 401 and \
                    time.monotonic() - token_refreshed > 120:
                # Tokens live ~1 hour; a bulk move run outlives them.
                # Re-acquire (same path the run started with) and retry this
                # message once. A 401 on a fresh token (permission not
                # granted) falls through as a failure.
                token = _write_token(config, ucfg)
                token_refreshed = time.monotonic()
                log_event(paths, "execute_token_refreshed", moved=total_moved)
                resp = _graph_post(
                    f"{GRAPH_API_BASE}/users/{mailbox}/messages/{mid}/move",
                    token, {"destinationId": dest})
            if resp is not None and resp.status_code in (200, 201):
                _append_jsonl(paths["moves"], {
                    "ts": _now_iso(), "id": mid, "cohort": cohort["id"],
                    "dest": dest, "source_folder": "inbox",
                })
                moved_ids.add(mid)
                done += 1
                total_moved += 1
                if done % 100 == 0:
                    log_event(paths, "cohort_progress", cohort=cohort["id"],
                              moved=done, of=len(ids))
            else:
                log_event(paths, "move_failed", cohort=cohort["id"], id=mid,
                          status=(resp.status_code if resp is not None else "none"))
        remaining = len(ids) - done
        if remaining == 0 and done > 0:
            cohort["status"] = "executed"
            cohort["executed_at"] = _now_iso()
        log_event(paths, "cohort_executed", cohort=cohort["id"], moved=done,
                  remaining=remaining)
        if limit is not None and total_moved >= limit:
            log_event(paths, "execute_limit_reached", limit=limit)
            break

    save_cohorts(paths, cohorts)
    log_event(paths, "execute_done", live=live, planned=total_planned,
              moved=total_moved)
    return {"planned": total_planned, "moved": total_moved}


# =============================================================================
# Status + CLI
# =============================================================================

def status(paths: dict) -> None:
    records = load_snapshot(paths)
    cohorts = load_cohorts(paths)
    state = load_state(paths)
    by_status: dict[str, int] = {}
    for c in cohorts:
        by_status[c.get("status", "?")] = by_status.get(c.get("status", "?"), 0) + 1
    print(f"inbox-{paths['user_key']} status")
    print(f"  snapshot:    {len(records):,} messages "
          f"(cursor {state.get('snapshot_cursor', 'none')})")
    print(f"  cohorts:     {len(cohorts)} " +
          ("(" + ", ".join(f"{v} {k}" for k, v in sorted(by_status.items())) + ")"
           if cohorts else ""))
    print(f"  moved:       {len(_load_moved_ids(paths)):,} messages")
    if state.get("q_answered_at"):
        q_txt = f"answered {state['q_answered_at'][:10]}"
    elif state.get("q_sent_at"):
        q_txt = f"emailed {state['q_sent_at'][:10]}, awaiting reply"
    else:
        q_txt = "not sent"
    print(f"  teams chat:  {'linked' if state.get('chat_id') else 'not yet created'}"
          f" | questionnaire {q_txt}")
    print(f"  rules file:  {paths['rules'] if paths['rules'].exists() else '(not created yet)'}")


def run_cli(user_key: str, config: dict, data_dir: Path) -> None:
    """CLI glue, called from rocky.py's dispatch for any --inbox-<user> flag."""
    from rocky import get_msal_app, acquire_app_token  # lazy (dispatch-time)

    users = config.get("inbox_users") or {}
    ucfg = users.get(user_key)
    if not ucfg or not ucfg.get("mailbox") or "PASTE" in ucfg.get("mailbox", ""):
        log.error(f"[inbox-{user_key}] no inbox_users['{user_key}'] entry with a "
                  f"real mailbox in config.json — see config.example.json.")
        sys.exit(1)

    paths = user_paths(config, data_dir, user_key)
    paths["share"].mkdir(parents=True, exist_ok=True)
    paths["local"].mkdir(parents=True, exist_ok=True)
    args = sys.argv[1:]

    def _limit() -> int | None:
        for i, a in enumerate(args):
            if a == "--limit" and i + 1 < len(args):
                try:
                    return int(args[i + 1])
                except ValueError:
                    return None
        return None

    def _hours(default: int) -> int:
        for i, a in enumerate(args):
            if a == "--hours" and i + 1 < len(args):
                try:
                    return max(1, int(args[i + 1]))
                except ValueError:
                    return default
        return default

    log.info("=" * 60)
    log.info(f"Rocky — Inbox Cleaner (inbox-{user_key}, "
             f"mailbox {ucfg['mailbox']})")
    log.info("=" * 60)

    if "--snapshot" in args:
        token = acquire_app_token(config)
        result = snapshot(token, config, ucfg, paths, full="--full" in args)
        log.info(f"[inbox-{user_key}] snapshot: {result}")
    elif "--analyze" in args:
        result = analyze(config, ucfg, paths)
        log.info(f"[inbox-{user_key}] analyze: {result}")
    elif "--questionnaire" in args:
        result = questionnaire_cycle(config, ucfg, paths,
                                     resend="--resend" in args)
        log.info(f"[inbox-{user_key}] questionnaire: {result}")
    elif "--chat" in args:
        result = chat_cycle(config, ucfg, paths)
        log.info(f"[inbox-{user_key}] chat cycle: {result}")
    elif "--rules-update" in args:
        result = rules_update(config, ucfg, paths)
        log.info(f"[inbox-{user_key}] rules update: {result}")
    elif "--digest" in args:
        result = digest_cycle(config, ucfg, paths, hours=_hours(24))
        log.info(f"[inbox-{user_key}] digest: {result}")
    elif "--internal-report" in args:
        result = internal_report(config, ucfg, paths)
        log.info(f"[inbox-{user_key}] internal report: {result}")
    elif "--execute" in args:
        live = "--live" in args
        token = _write_token(config, ucfg)
        result = execute(token, config, ucfg, paths, live=live, limit=_limit())
        mode = "LIVE" if live else "DRY-RUN"
        log.info(f"[inbox-{user_key}] execute ({mode}): {result}")
    elif "--status" in args:
        status(paths)
    else:
        print(f"inbox-{user_key} subcommands: --snapshot [--full] | --analyze | "
              f"--questionnaire [--resend] | --chat | --rules-update | "
              f"--digest [--hours N] | --internal-report | "
              f"--execute [--live] [--limit N] | --status")
        sys.exit(0)
