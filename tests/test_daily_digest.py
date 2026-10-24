"""Daily case digest: shared-RRID folders, docx tables, the spine's deadlines,
facts from filed pleadings, the scheduling-info list, and the email fetch
window. Claude is faked.

    python -m pytest tests/test_daily_digest.py -q
"""

import io
import json
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import rocky  # noqa: E402

TODAY = date(2026, 10, 19)


class FakeClient:
    """Returns canned text; records every prompt it was sent."""

    def __init__(self, reply):
        self.reply = reply
        self.prompts = []
        self.messages = self

    def create(self, **kw):
        self.prompts.append((kw.get("system", ""), kw["messages"][0]["content"]))
        text = self.reply(kw) if callable(self.reply) else self.reply
        return SimpleNamespace(content=[SimpleNamespace(text=text)])


def _root(tmp_path, monkeypatch, *names):
    for n in names:
        (tmp_path / n).mkdir()
    monkeypatch.setattr(rocky, "ROCKY_CASES_ROOT", tmp_path)
    return tmp_path


# --- Shared RRIDs ----------------------------------------------------------

def test_shared_rrid_picks_the_folder_matching_the_index_name(tmp_path, monkeypatch):
    _root(tmp_path, monkeypatch, "Board, Bianca (RRID-0040)",
          "Murray, Sheena (RRID-0040)", "Wall, Jesse (RRID-0040)", "Ricks, Sharday (RRID-0004)")
    assert rocky.find_case_folder("RRID-0040", "Wall, Jesse").name == "Wall, Jesse (RRID-0040)"
    # No name, or a name that matches none of them: refuse to guess.
    assert rocky.find_case_folder("RRID-0040") is None
    assert rocky.find_case_folder("RRID-0040", "Smith, Al") is None
    # A unique RRID needs no name, and a trailing space in the index is fine.
    assert rocky.find_case_folder("RRID-0004").name.startswith("Ricks")
    assert rocky.find_case_folder("RRID-0004", "Ricks, Sharday ").name.startswith("Ricks")
    assert rocky.find_rrid_collisions() == {"RRID-0040": [
        "Board, Bianca (RRID-0040)", "Murray, Sheena (RRID-0040)", "Wall, Jesse (RRID-0040)"]}


def test_rrid_match_is_exact_not_a_substring(tmp_path, monkeypatch):
    _root(tmp_path, monkeypatch, "Old (RRID-00401)", "New (RRID-0040)")
    assert rocky.find_case_folder("RRID-0040").name == "New (RRID-0040)"


# --- Word tables -----------------------------------------------------------

def test_docx_text_includes_tables_in_order():
    from docx import Document

    doc = Document()
    doc.add_paragraph("MASTER CASE SUMMARY")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Date", "Event"
    table.cell(1, 0).text, table.cell(1, 1).text = "2026-10-23", "Remote scheduling conference"
    doc.add_paragraph("After the table")
    buf = io.BytesIO()
    doc.save(buf)
    text = rocky._extract_docx(buf.getvalue(), "m.docx")
    assert text.splitlines() == [
        "MASTER CASE SUMMARY", "Date | Event",
        "2026-10-23 | Remote scheduling conference", "After the table"]


# --- CASE_SPINE deadlines --------------------------------------------------

SPINE = """# CASE SPINE

## Deadlines

| Date | Description | Source | Status |
|------|-------------|--------|--------|
| 2026-07-24 | Initial scheduling conference | PLD-014 | MISSED |
| TBD | Serve interrogatory answers | WP-005 | OPEN |
| TBD | Old task | X | DONE — finished |
| 2026-10-23 | Remote Scheduling Conference at 10:30 AM | PLD-021 | PENDING |

---

## Session Log
| 2027-01-01 | not a deadline row | | |
"""


def test_spine_deadlines_keep_open_and_upcoming_rows(tmp_path):
    (tmp_path / "CASE_SPINE.md").write_text(SPINE, encoding="utf-8")
    rows, upcoming = rocky._read_spine_deadlines(tmp_path, TODAY)
    assert "Serve interrogatory answers" in rows
    assert "2026-10-23" in rows
    assert "2026-07-24" not in rows and "Old task" not in rows and "2027-01-01" not in rows
    assert upcoming == [date(2026, 10, 23)]


def test_spine_missing_is_empty(tmp_path):
    assert rocky._read_spine_deadlines(tmp_path, TODAY) == ("", [])


# --- Facts from filed pleadings --------------------------------------------

def _facts_event(filed_by, facts, document="Opposition"):
    return {"timestamp": "2026-10-19T20:00:00Z", "event": "pleading_facts",
            "document": document, "filed_by": filed_by, "filed_on": "2026-10-19",
            "facts": [{"date": d, "fact": f} for d, f in facts]}


def test_pleading_facts_are_sorted_tagged_and_dated(tmp_path):
    lines = [
        _facts_event("us", [("2026-10-05", "Bozzuto cancelled the depositions."),
                            ("2026-08-26", "Bozzuto served deposition notices.")]),
        _facts_event("opposing", [("2026-11-30", "Plaintiff says trial is set.")], "Motion"),
        _facts_event("court", [("2026-10-30", "Court set a status hearing.")], "Order"),
        {"timestamp": "2026-10-19T20:00:00Z", "event": "research", "summary": "x"},
    ]
    (tmp_path / "activity.jsonl").write_text(
        "\n".join(json.dumps(e) for e in lines), encoding="utf-8")
    text, upcoming = rocky._read_pleading_facts(tmp_path, TODAY)
    out = text.splitlines()
    assert out[0].startswith("- 2026-08-26") and "[our filing: Opposition, filed 2026-10-19]" in out[0]
    assert "cancelled the depositions" in out[1]
    assert "[court: Order" in out[2]
    assert "[opposing party's allegation: Motion" in out[3]
    # The other side's dates don't count as scheduling info.
    assert upcoming == [date(2026, 10, 30)]


def test_log_pleading_facts_writes_event_and_claude_md(tmp_path):
    (tmp_path / "CLAUDE.md").write_text("# Case\n", encoding="utf-8")
    reply = json.dumps({"document": "Opposition", "filed_by": "us", "filed_on": "2026-10-19",
                        "facts": [{"date": "2026-10-05", "fact": "Depositions cancelled."},
                                  {"date": "", "fact": "undated, dropped"}]})
    ev = rocky.log_pleading_facts(FakeClient(reply), tmp_path, "RRID-0019", "Ghumman",
                                  "01 - Pleadings/PLD-090 Opp.pdf", "text of the filing")
    assert ev["facts"] == [{"date": "2026-10-05", "fact": "Depositions cancelled."}]
    logged = [json.loads(l) for l in (tmp_path / "activity.jsonl").read_text().splitlines()]
    assert logged[0]["event"] == "pleading_facts" and logged[0]["source"].endswith("Opp.pdf")
    assert "## Facts From Filed Pleadings" in (tmp_path / "CLAUDE.md").read_text(encoding="utf-8")
    # On its own it isn't "activity" — the document_filed event is.
    assert not rocky._is_substantive_event(logged[0])


def test_pleading_facts_prompt_names_our_client(tmp_path):
    # Board's folder is named for the tenant; the model must be told we
    # represent Bozzuto, not read "Board, Bianca" as our side.
    reply = json.dumps({"document": "Plaintiff's Opposition", "filed_by": "opposing",
                        "filed_on": "2026-10-12",
                        "facts": [{"date": "2026-10-12", "fact": "Opposition filed."}]})
    fake = FakeClient(reply)
    ev = rocky.log_pleading_facts(
        fake, tmp_path, "RRID-0041",
        "Board, Bianca — Client: Bozzuto Management Company. Wrongful eviction claim.",
        "Pleadings/PLD-009 Plaintiff's Opposition.pdf", "text of the filing",
        our_client="Bozzuto Management Company", case_name="Board, Bianca")
    assert ev["filed_by"] == "opposing"
    system, user = fake.prompts[0]
    assert ("WE REPRESENT: Bozzuto Management Company (Gallagher LLP / James Bragdon)."
            in user)
    assert "The case name 'Board, Bianca' may be the opposing party's name" in user
    assert system == rocky.PLEADING_FACTS_SYSTEM_PROMPT
    assert "client named in WE REPRESENT" in system
    assert "including the person the case is named after" in system
    assert 'If the filer is unclear, use "unknown", not "us"' in system
    assert "client named in CASE" not in system


def test_pleading_facts_prompt_without_client():
    fake = FakeClient(json.dumps({"filed_by": "unknown", "facts": []}))
    rocky.extract_pleading_facts(fake, "Wall, Jesse", "PLD-001.pdf", "text",
                                 our_client="unknown", case_name="Wall, Jesse")
    user = fake.prompts[0][1]
    assert "WE REPRESENT: unknown (Gallagher LLP / James Bragdon). Use filed_by" in user


def test_court_filing_folders():
    assert rocky._is_court_filing_folder("01 - Pleadings")
    assert rocky._is_court_filing_folder("02 - Motions")
    assert not rocky._is_court_filing_folder("05 - Email Correspondence")
    assert not rocky._is_court_filing_folder("04 - Fact Research")


# --- Scheduling info -------------------------------------------------------

def test_scheduling_status():
    with_notice = "- 01 - Pleadings/PLD-021 Notice of Remote Scheduling Conference.pdf"
    bare = "- 01 - Pleadings/PLD-001 Complaint.pdf"
    assert rocky._scheduling_status(bare, [date(2026, 10, 23)]) == "current"
    assert rocky._scheduling_status(with_notice, []) == "stale"
    assert rocky._scheduling_status(bare, []) == "none"
    assert rocky._scheduling_status("- Drafts/board minutes.docx", []) == ""


def test_past_next_deadline_is_dropped():
    assert rocky._drop_past_deadline("2026-08-26 — Status hearing (PAST)", TODAY) == "None on file"
    assert rocky._drop_past_deadline("2026-12-02 — Trial", TODAY) == "2026-12-02 — Trial"
    assert rocky._drop_past_deadline("None on file", TODAY) == "None on file"


# --- End to end --------------------------------------------------------------

def test_digest_skips_borrowed_rrid_flags_alerts_and_scheduling(tmp_path, monkeypatch):
    root = _root(tmp_path, monkeypatch, "Board, Bianca (RRID-0040)",
                 "Wall, Jesse (RRID-0040)", "Ricks, Sharday (RRID-0004)",
                 "Ghumman, Saira (RRID-0019)")
    for case in ("Board, Bianca (RRID-0040)", "Wall, Jesse (RRID-0040)",
                 "Ricks, Sharday (RRID-0004)", "Ghumman, Saira (RRID-0019)"):
        (root / case / "01 - Pleadings").mkdir()
        (root / case / "01 - Pleadings" / "PLD-001 Complaint.pdf").write_text("x")
    (root / "Ricks, Sharday (RRID-0004)" / "CASE_SPINE.md").write_text(SPINE, encoding="utf-8")
    now = "2026-10-19T20:00:00Z"
    (root / "Ghumman, Saira (RRID-0019)" / "activity.jsonl").write_text(
        json.dumps({"timestamp": now, "actor": "Sang (Cowork)", "event": "drafting",
                    "summary": "Revised the opposition."}) + "\n"
        + json.dumps(_facts_event("us", [("2026-10-05", "Bozzuto cancelled the depositions.")])),
        encoding="utf-8")
    monkeypatch.setattr(rocky, "DAILY_DIGESTS_DIR", root / "Daily Digests")
    monkeypatch.setattr(rocky, "load_case_index", lambda: [
        {"RRID#": "RRID-0040", "File Name": "Wall, Jesse"},
        {"RRID#": "RRID-0004", "File Name": "Ricks, Sharday "},
        {"RRID#": "RRID-0019", "File Name": "Ghumman, Saira", "Case Folder": "Inbox\\Ghumman"},
    ])
    monkeypatch.setattr(rocky, "datetime", _FrozenDatetime)

    def reply(kw):
        if "BIG DEADLINE" in kw["system"]:
            return ("RRID-0040 | Coordinate with eviction counsel | None on file\n"
                    "RRID-0004 | Prepare for conference | 2026-10-23 — Remote Scheduling Conference")
        return "**What happened**\n- Opposition revised.\n\n**Upcoming dates**\n(none on file)"

    client = FakeClient(reply)
    result = rocky.daily_digest(client, "")
    md = (root / "Daily Digests" / "2026-10-19.md").read_text(encoding="utf-8")

    assert md.count("**RRID-0040 — Wall, Jesse** — No new activity") == 1
    assert "Board, Bianca" in md.split("## Rocky Needs Attention")[1].split("---")[0]
    assert "RRID-0040 is on 2 folders" in md
    assert "Email not monitored" in md and "RRID-0004 Ricks, Sharday" in md
    sched = md.split("## Scheduling Info")[1].split("## Cases with No Activity")[0]
    assert "RRID-0004 — Ricks, Sharday: 2026-10-23" in sched.split("**On file**")[1]
    assert "RRID-0040 — Wall, Jesse: no scheduling order" in sched
    assert "RRID-0019 — Ghumman, Saira: no scheduling order" in sched
    # The Ghumman section saw our filing's facts and the trust rules.
    system, prompt = next(p for p in client.prompts if "FACTS FROM FILED PLEADINGS (" in p[1])
    assert "Bozzuto cancelled the depositions. [our filing: Opposition" in prompt
    assert "do not ask whether the depositions went forward" in system
    # The no-activity call got Ricks's deadlines table ahead of the memo.
    _, batch = next(p for p in client.prompts if "BIG DEADLINE" in p[0])
    assert batch.index("DEADLINES TABLE:") < batch.index("STATUS MEMO:")
    assert result["written"]


class _FrozenDatetime(datetime):
    @classmethod
    def now(cls, tz=None):
        return datetime(2026, 10, 19, 21, 0, tzinfo=timezone.utc)


# --- Email fetch window ----------------------------------------------------

def test_fetch_window_starts_at_last_completed_run(tmp_path, monkeypatch):
    mark = tmp_path / "daily_cases_last_run.json"
    monkeypatch.setattr(rocky, "DAILY_CASES_WATERMARK", mark)
    monkeypatch.setattr(rocky, "STATE_DIR", tmp_path)
    now = datetime(2026, 10, 20, 20, 0, tzinfo=timezone.utc)
    assert rocky._daily_cases_since(now) == now - timedelta(hours=24)
    rocky._save_daily_cases_watermark(now - timedelta(hours=48))
    assert rocky._daily_cases_since(now) == now - timedelta(hours=48)
    rocky._save_daily_cases_watermark(now - timedelta(days=30))
    assert rocky._daily_cases_since(now) == now - timedelta(hours=24)


def test_daily_cases_without_case_folder_column_keeps_watermark(tmp_path, monkeypatch):
    _root(tmp_path, monkeypatch, "Castillo, Jose (RRID-0036)")
    mark = tmp_path / "mark.json"
    monkeypatch.setattr(rocky, "DAILY_CASES_WATERMARK", mark)
    monkeypatch.setattr(rocky, "load_case_index",
                        lambda: [{"RRID#": "RRID-0036", "File Name": "Castillo, Jose "}])
    results = rocky.daily_cases(FakeClient("{}"), "tok", "j@x", "")
    assert results[0]["reason"] == "no_case_folder"
    assert not mark.exists()
    assert rocky.open_cases_without_mail_folder(
        [{"RRID#": "RRID-0036"}, {"RRID#": "RRID-0001", "Case Folder": "Inbox\\X"},
         {"RRID#": "RRID-0002", "Open/Closed": "Closed"}]) == [{"RRID#": "RRID-0036"}]


def test_section_labels_normalized_to_bold():
    body = "## What happened\n- x\n\n### **Upcoming dates**\n(none on file)\n- ## not a label"
    assert rocky._normalize_section_labels(body) == (
        "**What happened**\n- x\n\n**Upcoming dates**\n(none on file)\n- ## not a label")
